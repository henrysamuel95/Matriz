# -*- coding: utf-8 -*-
import base64
import io
from datetime import date, datetime, time

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

from odoo import _, fields, models
from odoo.exceptions import UserError


class ActaEntregaInformeLinea(models.TransientModel):
    _name = 'acta.entrega.informe.linea'
    _description = 'Línea de informe de cuentas de cobro'
    _order = 'ultima_fecha_envio desc, fecha_creacion desc, id desc'

    wizard_id = fields.Many2one(
        'acta.entrega.informe.wizard',
        string='Informe',
        required=True,
        ondelete='cascade',
        index=True,
    )
    acta_id = fields.Integer(string='ID cuenta', readonly=True)
    contratista = fields.Char(string='Contratista', readonly=True)
    documento = fields.Char(string='Documento/NIT', readonly=True)
    tipo_persona = fields.Char(string='Tipo de persona', readonly=True)
    proyecto = fields.Char(string='Proyecto', readonly=True)
    tipo_vinculacion = fields.Char(string='Tipo', readonly=True)
    origen_cuenta = fields.Char(string='Origen', readonly=True)
    estado = fields.Char(string='Estado actual', readonly=True)
    estado_visible = fields.Char(string='Estado visible', readonly=True)
    pago = fields.Integer(string='Pago N°', readonly=True)
    numero_factura = fields.Char(string='Factura', readonly=True)
    valor_cuenta = fields.Float(string='Valor cuenta', readonly=True)

    fecha_creacion = fields.Datetime(string='Fecha creación', readonly=True)
    fecha_radicacion = fields.Datetime(string='Fecha radicación', readonly=True)
    fecha_estado_actual = fields.Datetime(string='Fecha último estado', readonly=True)
    ultima_fecha_envio = fields.Datetime(string='Último envío/movimiento', readonly=True)
    ultimo_movimiento = fields.Char(string='Último movimiento', readonly=True)

    fecha_asistente = fields.Datetime(string='Asistente coordinación', readonly=True)
    fecha_coordinador = fields.Datetime(string='Coordinador', readonly=True)
    fecha_ventanilla = fields.Datetime(string='Ventanilla', readonly=True)
    fecha_director_financiero = fields.Datetime(string='Director Financiero', readonly=True)
    fecha_firma_director_financiero = fields.Date(string='Firma Director Financiero', readonly=True)
    fecha_contabilidad = fields.Datetime(string='Contabilidad', readonly=True)
    fecha_aux_tesoreria = fields.Datetime(string='Auxiliar tesorería', readonly=True)
    fecha_tesoreria = fields.Datetime(string='Tesorería', readonly=True)
    fecha_finalizacion = fields.Datetime(string='Finalización', readonly=True)
    fecha_rechazo = fields.Datetime(string='Rechazo', readonly=True)

    horas_creacion_asistente = fields.Float(string='Horas creación → Asistente', readonly=True, group_operator='avg')
    horas_asistente_coordinador = fields.Float(string='Horas Asistente → Coordinador', readonly=True, group_operator='avg')
    horas_coordinador_ventanilla = fields.Float(string='Horas Coordinador → Ventanilla', readonly=True, group_operator='avg')
    horas_ventanilla_director = fields.Float(string='Horas Ventanilla → Director Financiero', readonly=True, group_operator='avg')
    horas_director_contabilidad = fields.Float(string='Horas Director Financiero → Contabilidad', readonly=True, group_operator='avg')
    horas_contabilidad_aux = fields.Float(string='Horas Contabilidad → Aux. tesorería', readonly=True, group_operator='avg')
    horas_aux_tesoreria = fields.Float(string='Horas Aux. tesorería → Tesorería', readonly=True, group_operator='avg')
    horas_tesoreria_final = fields.Float(string='Horas Tesorería → Finalizado', readonly=True, group_operator='avg')
    duracion_total_horas = fields.Float(string='Duración total (horas)', readonly=True, group_operator='avg')
    duracion_total_dias = fields.Float(string='Duración total (días)', readonly=True, group_operator='avg')
    horas_estado_actual = fields.Float(string='Horas en estado actual', readonly=True, group_operator='avg')
    duracion_correccion_horas = fields.Float(string='Duración corrección (horas)', readonly=True, group_operator='avg')

    en_correccion = fields.Boolean(string='En corrección', readonly=True)
    seguimiento_correccion = fields.Char(string='Seguimiento corrección', readonly=True)
    cesion_derechos = fields.Boolean(string='Cesión de derechos', readonly=True)
    estado_cesion = fields.Char(string='Estado cesión', readonly=True)
    finalizada = fields.Boolean(string='Finalizada', readonly=True)
    rechazada = fields.Boolean(string='Rechazada', readonly=True)


class ActaEntregaInformeWizard(models.TransientModel):
    _name = 'acta.entrega.informe.wizard'
    _description = 'Generador de informes de cuentas de cobro'

    TIPO_INFORME = [
        ('general', 'General - todas las cuentas'),
        ('fechas', 'Por rango de fechas'),
        ('ultimos', 'Últimos movimientos / enviados'),
        ('tiempos', 'Tiempos entre estados'),
        ('correcciones', 'Correcciones'),
        ('finalizadas', 'Finalizadas'),
        ('rechazadas', 'Rechazadas'),
        ('cesiones', 'Cesiones de derechos'),
    ]

    CAMPO_FECHA = [
        ('create_date', 'Fecha de creación'),
        ('fecha_solicitud', 'Fecha de radicación'),
        ('fecha_estado_actual', 'Fecha del último estado'),
        ('fecha_envio_asis_coordinador', 'Fecha envío Asistente coordinación'),
        ('fecha_envio_coordinador', 'Fecha envío Coordinador'),
        ('fecha_envio_ventanilla', 'Fecha envío Ventanilla'),
        ('fecha_envio_director_financiero', 'Fecha envío Director Financiero'),
        ('fecha_envio_contabilidad', 'Fecha envío Contabilidad'),
        ('fecha_envio_auxiliar_tesoreria', 'Fecha envío Auxiliar tesorería'),
        ('fecha_envio_tesoreria', 'Fecha envío Tesorería'),
        ('fecha_finalizacion', 'Fecha finalización'),
        ('fecha_rechazo', 'Fecha rechazo'),
    ]

    ESTADOS = [
        ('asis_coordinador', 'Asistente de coordinación'),
        ('coordinador', 'Coordinador'),
        ('ventanilla', 'Ventanilla'),
        ('director_financiero', 'Director Financiero'),
        ('finaciera', 'Contabilidad'),
        ('au_tesoreria', 'Auxiliar de tesorería'),
        ('tesorería', 'Tesorería'),
        ('finalizado', 'Finalizado'),
        ('rechazado', 'Rechazado'),
    ]

    tipo_informe = fields.Selection(TIPO_INFORME, string='Tipo de informe', default='general', required=True)
    campo_fecha = fields.Selection(CAMPO_FECHA, string='Fecha a consultar', default='create_date', required=True)
    fecha_desde = fields.Datetime(string='Desde')
    fecha_hasta = fields.Datetime(string='Hasta')
    estado = fields.Selection(ESTADOS, string='Estado')
    tipo_persona = fields.Selection([
        ('persona_natural', 'Persona natural'),
        ('persona_juridica', 'Persona jurídica'),
        ('seleccion_manual', 'Selección manual'),
    ], string='Tipo de persona')
    tipo_vinculacion = fields.Selection([
        ('contrato', 'Contrato'),
        ('otrosi', 'Otrosí'),
        ('bienes', 'Bienes'),
        ('servicios', 'Servicios'),
        ('contrato_laboral', 'Contrato laboral'),
    ], string='Tipo de vinculación')
    solo_correcciones = fields.Boolean(string='Solo cuentas con corrección activa')
    solo_cesiones = fields.Boolean(string='Solo cuentas con cesión de derechos')
    limite = fields.Integer(
        string='Máximo de registros',
        default=0,
        help='0 significa todas las cuentas que cumplan los filtros.',
    )
    ordenar_por = fields.Selection([
        ('ultima_desc', 'Más reciente primero'),
        ('ultima_asc', 'Más antiguo primero'),
        ('demora_desc', 'Mayor demora primero'),
        ('demora_asc', 'Menor demora primero'),
    ], string='Ordenar por', default='ultima_desc', required=True)

    archivo_xlsx = fields.Binary(string='Archivo XLSX', readonly=True, attachment=False)
    archivo_nombre = fields.Char(string='Nombre archivo', readonly=True)
    total_resultados = fields.Integer(string='Resultados', readonly=True)

    @staticmethod
    def _to_datetime(value):
        if not value:
            return False
        if isinstance(value, datetime):
            return value
        if isinstance(value, date):
            return datetime.combine(value, time.min)
        try:
            return fields.Datetime.to_datetime(value)
        except Exception:
            return False

    @classmethod
    def _horas(cls, inicio, fin):
        inicio = cls._to_datetime(inicio)
        fin = cls._to_datetime(fin)
        if not inicio or not fin:
            return 0.0
        return round(max(0.0, (fin - inicio).total_seconds() / 3600.0), 2)

    @staticmethod
    def _selection_label(record, field_name):
        if field_name not in record._fields:
            return ''
        field = record._fields[field_name]
        value = record[field_name]
        if not value:
            return ''
        selection = field.selection(record.env) if callable(field.selection) else field.selection
        return dict(selection or []).get(value, value)

    def _domain_actas(self):
        self.ensure_one()
        Acta = self.env['acta_entrega_documentos']
        domain = []

        if self.estado:
            domain.append(('estado', '=', self.estado))
        if self.tipo_persona and 'clase_persona' in Acta._fields:
            domain.append(('clase_persona', '=', self.tipo_persona))
        if self.tipo_vinculacion and 'selecionar' in Acta._fields:
            domain.append(('selecionar', '=', self.tipo_vinculacion))
        if self.solo_correcciones and 'correccion_portal_activa' in Acta._fields:
            domain.append(('correccion_portal_activa', '=', True))
        if self.solo_cesiones and 'cesion_derechos' in Acta._fields:
            domain.append(('cesion_derechos', '=', True))

        if self.tipo_informe == 'correcciones' and 'correccion_portal_activa' in Acta._fields:
            domain.append(('correccion_portal_activa', '=', True))
        elif self.tipo_informe == 'finalizadas':
            domain.append(('estado', '=', 'finalizado'))
        elif self.tipo_informe == 'rechazadas':
            domain.append(('estado', '=', 'rechazado'))
        elif self.tipo_informe == 'cesiones' and 'cesion_derechos' in Acta._fields:
            domain.append(('cesion_derechos', '=', True))

        campo_fecha = self.campo_fecha or 'create_date'
        if campo_fecha not in Acta._fields:
            raise UserError(_('El campo de fecha seleccionado no existe en cuentas de cobro.'))
        if self.fecha_desde:
            domain.append((campo_fecha, '>=', self.fecha_desde))
        if self.fecha_hasta:
            domain.append((campo_fecha, '<=', self.fecha_hasta))

        return domain

    def _valor_fecha(self, record, field_name):
        if field_name not in record._fields:
            return False
        return self._to_datetime(record[field_name])

    def _line_vals(self, record):
        fecha_creacion = self._valor_fecha(record, 'fch_creacion') or self._valor_fecha(record, 'create_date')
        fecha_radicacion = self._valor_fecha(record, 'fecha_solicitud')
        fecha_estado_actual = self._valor_fecha(record, 'fecha_estado_actual')
        fecha_asistente = self._valor_fecha(record, 'fecha_envio_asis_coordinador')
        fecha_coordinador = self._valor_fecha(record, 'fecha_envio_coordinador')
        fecha_ventanilla = self._valor_fecha(record, 'fecha_envio_ventanilla')
        fecha_director = self._valor_fecha(record, 'fecha_envio_director_financiero')
        fecha_contabilidad = self._valor_fecha(record, 'fecha_envio_contabilidad')
        fecha_aux = self._valor_fecha(record, 'fecha_envio_auxiliar_tesoreria')
        fecha_tesoreria = self._valor_fecha(record, 'fecha_envio_tesoreria')
        fecha_finalizacion = self._valor_fecha(record, 'fecha_finalizacion')
        fecha_rechazo = self._valor_fecha(record, 'fecha_rechazo')

        movimientos = [
            (fecha_asistente, 'Envío a Asistente de coordinación'),
            (fecha_coordinador, 'Envío a Coordinador'),
            (fecha_ventanilla, 'Envío a Ventanilla'),
            (fecha_director, 'Envío a Director Financiero'),
            (fecha_contabilidad, 'Envío a Contabilidad'),
            (fecha_aux, 'Envío a Auxiliar de tesorería'),
            (fecha_tesoreria, 'Envío a Tesorería'),
            (fecha_finalizacion, 'Finalización'),
            (fecha_rechazo, 'Rechazo'),
        ]
        movimientos = [(fecha, nombre) for fecha, nombre in movimientos if fecha]
        if movimientos:
            ultima_fecha, ultimo_movimiento = max(movimientos, key=lambda item: item[0])
        else:
            ultima_fecha = fecha_estado_actual or fecha_radicacion or fecha_creacion
            ultimo_movimiento = 'Creación / estado actual'

        inicio_total = fecha_radicacion or fecha_creacion
        fin_total = fecha_finalizacion or fecha_rechazo or ultima_fecha or fecha_estado_actual
        horas_total = self._horas(inicio_total, fin_total)

        estado_tecnico = record.estado if 'estado' in record._fields else False
        horas_estado_actual = 0.0
        if fecha_estado_actual and estado_tecnico not in ('finalizado', 'rechazado'):
            horas_estado_actual = self._horas(fecha_estado_actual, fields.Datetime.now())

        corr_ini = self._valor_fecha(record, 'correccion_reporte_fecha_solicitud')
        corr_fin = self._valor_fecha(record, 'correccion_reporte_fecha_cierre')
        corr_activa = bool(getattr(record, 'correccion_portal_activa', False))
        duracion_correccion = self._horas(corr_ini, corr_fin or (fields.Datetime.now() if corr_activa else False))

        proyecto = ''
        if 'proyecto_relacionado' in record._fields and record.proyecto_relacionado:
            proyecto = record.proyecto_relacionado.sudo().display_name or ''
        elif 'proyecto_nombre' in record._fields:
            proyecto = record.proyecto_nombre or ''

        return {
            'wizard_id': self.id,
            'acta_id': record.id,
            'contratista': getattr(record, 'nombre_contratista', False) or getattr(record, 'name', False) or '',
            'documento': getattr(record, 'documento_de_identificacion', False) or '',
            'tipo_persona': self._selection_label(record, 'clase_persona'),
            'proyecto': proyecto,
            'tipo_vinculacion': self._selection_label(record, 'selecionar'),
            'origen_cuenta': self._selection_label(record, 'origen_cuenta'),
            'estado': self._selection_label(record, 'estado'),
            'estado_visible': getattr(record, 'estado_kanban_display', False) or self._selection_label(record, 'estado'),
            'pago': int(getattr(record, 'pago', 0) or 0),
            'numero_factura': getattr(record, 'numero_factura', False) or '',
            'valor_cuenta': float(getattr(record, 'valor_cuenta', 0.0) or 0.0),
            'fecha_creacion': fecha_creacion,
            'fecha_radicacion': fecha_radicacion,
            'fecha_estado_actual': fecha_estado_actual,
            'ultima_fecha_envio': ultima_fecha,
            'ultimo_movimiento': ultimo_movimiento,
            'fecha_asistente': fecha_asistente,
            'fecha_coordinador': fecha_coordinador,
            'fecha_ventanilla': fecha_ventanilla,
            'fecha_director_financiero': fecha_director,
            'fecha_firma_director_financiero': getattr(record, 'fecha_firma_director_financiero', False) or False,
            'fecha_contabilidad': fecha_contabilidad,
            'fecha_aux_tesoreria': fecha_aux,
            'fecha_tesoreria': fecha_tesoreria,
            'fecha_finalizacion': fecha_finalizacion,
            'fecha_rechazo': fecha_rechazo,
            'horas_creacion_asistente': self._horas(inicio_total, fecha_asistente),
            'horas_asistente_coordinador': self._horas(fecha_asistente, fecha_coordinador),
            'horas_coordinador_ventanilla': self._horas(fecha_coordinador, fecha_ventanilla),
            'horas_ventanilla_director': self._horas(fecha_ventanilla, fecha_director),
            'horas_director_contabilidad': self._horas(fecha_director, fecha_contabilidad),
            'horas_contabilidad_aux': self._horas(fecha_contabilidad, fecha_aux),
            'horas_aux_tesoreria': self._horas(fecha_aux, fecha_tesoreria),
            'horas_tesoreria_final': self._horas(fecha_tesoreria, fecha_finalizacion),
            'duracion_total_horas': horas_total,
            'duracion_total_dias': round(horas_total / 24.0, 2) if horas_total else 0.0,
            'horas_estado_actual': horas_estado_actual,
            'duracion_correccion_horas': duracion_correccion,
            'en_correccion': corr_activa,
            'seguimiento_correccion': getattr(record, 'seguimiento_correccion_reporte', False) or '',
            'cesion_derechos': bool(getattr(record, 'cesion_derechos', False)),
            'estado_cesion': self._selection_label(record, 'estado_cesion_derechos'),
            'finalizada': estado_tecnico == 'finalizado',
            'rechazada': estado_tecnico == 'rechazado',
        }

    def _generar_lineas(self):
        self.ensure_one()
        Linea = self.env['acta.entrega.informe.linea'].sudo()
        Linea.search([('wizard_id', '=', self.id)]).unlink()

        actas = self.env['acta_entrega_documentos'].sudo().search(self._domain_actas())
        vals_list = [self._line_vals(record) for record in actas]

        if self.ordenar_por == 'ultima_desc':
            vals_list.sort(key=lambda x: x.get('ultima_fecha_envio') or datetime.min, reverse=True)
        elif self.ordenar_por == 'ultima_asc':
            vals_list.sort(key=lambda x: x.get('ultima_fecha_envio') or datetime.max)
        elif self.ordenar_por == 'demora_desc':
            vals_list.sort(key=lambda x: x.get('duracion_total_horas') or 0.0, reverse=True)
        elif self.ordenar_por == 'demora_asc':
            vals_list.sort(key=lambda x: x.get('duracion_total_horas') or 0.0)

        if self.limite and self.limite > 0:
            vals_list = vals_list[:self.limite]

        if vals_list:
            Linea.create(vals_list)
        self.total_resultados = len(vals_list)
        return Linea.search([('wizard_id', '=', self.id)])

    def action_generar(self):
        self.ensure_one()
        lineas = self._generar_lineas()
        action = self.env.ref('Matriz.action_acta_entrega_informe_lineas').sudo().read()[0]
        action['domain'] = [('id', 'in', lineas.ids)]
        action['context'] = {
            'search_default_group_estado': 0,
            'create': False,
            'edit': False,
            'delete': False,
        }
        return action

    def action_descargar_xlsx(self):
        self.ensure_one()
        lineas = self._generar_lineas()
        if not lineas:
            raise UserError(_('No hay cuentas de cobro que cumplan los filtros seleccionados.'))

        wb = Workbook()
        ws_resumen = wb.active
        ws_resumen.title = 'Resumen'
        ws_detalle = wb.create_sheet('Detalle')

        estados = {}
        for linea in lineas:
            estados[linea.estado or 'Sin estado'] = estados.get(linea.estado or 'Sin estado', 0) + 1

        resumen = [
            ('Informe', dict(self.TIPO_INFORME).get(self.tipo_informe, self.tipo_informe)),
            ('Total cuentas', len(lineas)),
            ('Finalizadas', sum(1 for x in lineas if x.finalizada)),
            ('Rechazadas', sum(1 for x in lineas if x.rechazada)),
            ('En corrección', sum(1 for x in lineas if x.en_correccion)),
            ('Con cesión de derechos', sum(1 for x in lineas if x.cesion_derechos)),
            ('Promedio duración total (horas)', round(sum(x.duracion_total_horas for x in lineas) / len(lineas), 2)),
            ('Promedio horas en estado actual', round(sum(x.horas_estado_actual for x in lineas) / len(lineas), 2)),
            ('Generado por', self.env.user.name),
            ('Fecha de generación', fields.Datetime.to_string(fields.Datetime.now())),
        ]
        ws_resumen.append(['Indicador', 'Valor'])
        for cell in ws_resumen[1]:
            cell.font = Font(bold=True)
        for item in resumen:
            ws_resumen.append(list(item))
        ws_resumen.append([])
        ws_resumen.append(['Estado', 'Cantidad'])
        for cell in ws_resumen[ws_resumen.max_row]:
            cell.font = Font(bold=True)
        for estado, cantidad in sorted(estados.items()):
            ws_resumen.append([estado, cantidad])
        ws_resumen.column_dimensions['A'].width = 38
        ws_resumen.column_dimensions['B'].width = 28

        columnas = [
            ('ID', 'acta_id'), ('Contratista', 'contratista'), ('Documento/NIT', 'documento'),
            ('Tipo persona', 'tipo_persona'), ('Proyecto', 'proyecto'), ('Tipo', 'tipo_vinculacion'),
            ('Origen', 'origen_cuenta'), ('Estado', 'estado'), ('Estado visible', 'estado_visible'),
            ('Pago N°', 'pago'), ('Factura', 'numero_factura'), ('Valor cuenta', 'valor_cuenta'),
            ('Fecha creación', 'fecha_creacion'), ('Fecha radicación', 'fecha_radicacion'),
            ('Último movimiento', 'ultimo_movimiento'), ('Último envío/movimiento', 'ultima_fecha_envio'),
            ('Fecha Asistente', 'fecha_asistente'), ('Fecha Coordinador', 'fecha_coordinador'),
            ('Fecha Ventanilla', 'fecha_ventanilla'), ('Fecha Director Financiero', 'fecha_director_financiero'),
            ('Fecha firma Director Financiero', 'fecha_firma_director_financiero'),
            ('Fecha Contabilidad', 'fecha_contabilidad'), ('Fecha Aux. tesorería', 'fecha_aux_tesoreria'),
            ('Fecha Tesorería', 'fecha_tesoreria'), ('Fecha finalización', 'fecha_finalizacion'),
            ('Fecha rechazo', 'fecha_rechazo'),
            ('Horas creación→Asistente', 'horas_creacion_asistente'),
            ('Horas Asistente→Coordinador', 'horas_asistente_coordinador'),
            ('Horas Coordinador→Ventanilla', 'horas_coordinador_ventanilla'),
            ('Horas Ventanilla→Director', 'horas_ventanilla_director'),
            ('Horas Director→Contabilidad', 'horas_director_contabilidad'),
            ('Horas Contabilidad→Aux.', 'horas_contabilidad_aux'),
            ('Horas Aux.→Tesorería', 'horas_aux_tesoreria'),
            ('Horas Tesorería→Finalizado', 'horas_tesoreria_final'),
            ('Duración total horas', 'duracion_total_horas'), ('Duración total días', 'duracion_total_dias'),
            ('Horas estado actual', 'horas_estado_actual'), ('Duración corrección horas', 'duracion_correccion_horas'),
            ('En corrección', 'en_correccion'), ('Seguimiento corrección', 'seguimiento_correccion'),
            ('Cesión derechos', 'cesion_derechos'), ('Estado cesión', 'estado_cesion'),
        ]
        ws_detalle.append([titulo for titulo, _campo in columnas])
        for cell in ws_detalle[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

        for linea in lineas:
            fila = []
            for _titulo, campo in columnas:
                valor = getattr(linea, campo)
                if isinstance(valor, datetime):
                    valor = valor.strftime('%Y-%m-%d %H:%M:%S')
                elif isinstance(valor, date):
                    valor = valor.strftime('%Y-%m-%d')
                elif isinstance(valor, bool):
                    valor = 'Sí' if valor else 'No'
                fila.append(valor if valor is not False else '')
            ws_detalle.append(fila)

        ws_detalle.freeze_panes = 'A2'
        ws_detalle.auto_filter.ref = ws_detalle.dimensions
        for idx, (titulo, _campo) in enumerate(columnas, start=1):
            ancho = min(max(len(titulo) + 2, 14), 34)
            ws_detalle.column_dimensions[get_column_letter(idx)].width = ancho

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        filename = 'informe_cuentas_cobro_%s.xlsx' % fields.Date.today().strftime('%Y%m%d')
        self.write({
            'archivo_xlsx': base64.b64encode(buffer.read()),
            'archivo_nombre': filename,
        })
        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/%s/%s/archivo_xlsx/%s?download=true' % (
                self._name,
                self.id,
                filename,
            ),
            'target': 'self',
        }
