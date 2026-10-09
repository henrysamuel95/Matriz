# -*- coding: utf-8 -*-
from odoo import models, fields
from odoo.exceptions import ValidationError
import base64
import io
import logging
import openai
import unicodedata

_logger = logging.getLogger(__name__)

try:
    import openpyxl
except ImportError:
    openpyxl = None
    _logger.warning("Falta openpyxl")


class SgoDestacadosPost(models.Model):
    _name = 'sgo.destacados.post'
    _description = 'Posts SGO Academy'
    _order = 'fecha_publicacion desc'

    name = fields.Char(string="Título", required=True)
    modelo_tecnico = fields.Char(string="Modelo técnico", index=True)
    descripcion_usuario = fields.Text(string="Descripción para usuarios")
    tipo = fields.Char(string="Tipo", index=True)
    descripcion_desarrollador = fields.Text(string="Descripción desarrollador")
    tiene_menu = fields.Boolean(string="Tiene menú principal")
    ruta = fields.Text(string="Ruta de acceso")
    detalles = fields.Text(string="Más detalles")
    contenido = fields.Html(string="Contenido")
    fecha_publicacion = fields.Datetime(string="Fecha Publicación", default=fields.Datetime.now)
    estado = fields.Selection([
        ('borrador', 'Borrador'),
        ('publicado', 'Publicado')
    ], string="Estado", default='borrador')
    pasos_ids = fields.One2many('sgo.guia.paso', 'post_id', string="Pasos del Tutorial")


class AgentePublicador(models.Model):
    _name = 'sgo.agente.publicador'
    _description = 'IA Generativa con OpenAI'

    name = fields.Char(default="Profesor de SGO Academy")
    archivo_datos = fields.Binary(string="Base de Conocimiento", required=True)
    archivo_nombre = fields.Char(string="Nombre Archivo")
    ultima_fila_procesada = fields.Integer(string="Última fila leída", default=1)
    api_key_ia = fields.Char(string="OpenAI API Key", help="Ingresa tu clave de OpenAI.")

    def _normalizar_cabecera(self, value):
        texto = str(value or '').strip().lower()
        texto = ''.join(
            char for char in unicodedata.normalize('NFD', texto)
            if unicodedata.category(char) != 'Mn'
        )
        return texto.replace('¿', '').replace('?', '')

    def _texto_excel(self, value):
        if value is None:
            return ''
        return str(value).strip()

    def _es_si(self, value):
        return self._normalizar_cabecera(value) in ('si', 'sí', 's')

    def accion_crear_publicacion(self):
        self.ensure_one()
        self.flush_recordset(['api_key_ia', 'archivo_datos'])

        if not self.archivo_datos:
            raise ValidationError("Sube el archivo Excel primero.")
        if not self.api_key_ia:
            raise ValidationError("Ingresa tu API Key de OpenAI.")
        if openpyxl is None:
            raise ValidationError("No está instalado openpyxl en el entorno de Odoo.")

        try:
            client = openai.OpenAI(api_key=self.api_key_ia)
            excel_data = base64.b64decode(self.archivo_datos)
            workbook = openpyxl.load_workbook(io.BytesIO(excel_data), data_only=True)
            sheet = workbook.active

            columnas = {
                self._normalizar_cabecera(sheet.cell(row=1, column=col).value): col
                for col in range(1, sheet.max_column + 1)
                if sheet.cell(row=1, column=col).value
            }

            fila_actual = self.ultima_fila_procesada + 1

            def get_val(*nombres):
                for nombre in nombres:
                    col = columnas.get(self._normalizar_cabecera(nombre))
                    if col:
                        return self._texto_excel(sheet.cell(row=fila_actual, column=col).value)
                return ''

            nombre_modelo = get_val('nombre del modelo')
            desc_usuario = get_val('descripción', 'descripcion')
            tipo = get_val('tipo')
            desc_dev = get_val('descripción desarrollador', 'descripcion desarrollador')
            tiene_menu_raw = get_val('tiene menú principal', 'tiene menu principal', 'tiene menú', 'tiene menu')
            ruta = get_val('ruta')
            detalles = get_val('más detalles', 'mas detalles')

            if not nombre_modelo:
                self.write({'ultima_fila_procesada': 1})
                return

            tiene_menu = self._es_si(tiene_menu_raw)
            ruta_texto = ruta or 'Uso interno del sistema'
            explicacion_base = desc_dev or desc_usuario or 'Componente del módulo Matriz.'

            prompt = f"""
Eres el profesor funcional de SGO Academy. Convierte esta ficha técnica en una explicación clara para usuarios internos.

Modelo técnico: {nombre_modelo}
Nombre funcional: {desc_usuario or nombre_modelo}
Tipo: {tipo or 'No clasificado'}
Tiene menú principal: {'Sí' if tiene_menu else 'No'}
Ruta de acceso: {ruta_texto}
Descripción técnica: {explicacion_base}
Más detalles: {detalles or 'Sin detalles adicionales'}

Entrega HTML breve y útil, sin markdown, con esta estructura exacta:
<b>¿Para qué sirve?</b><br/>Explica en lenguaje simple el propósito real.
<br/><br/><b>¿Dónde lo encuentro?</b><br/>Indica la ruta si existe; si es interno, di desde qué proceso o registro se activa.
<br/><br/><b>Cómo explicarlo al equipo</b><br/>Una frase práctica para que soporte lo comunique bien.

Máximo 110 palabras. No inventes rutas ni permisos. No uses emojis.
"""

            response = client.chat.completions.create(
                model="gpt-4o-mini",
                max_tokens=260,
                temperature=0.35,
                messages=[{"role": "user", "content": prompt}]
            )

            contenido = response.choices[0].message.content

            self.env['sgo.destacados.post'].create({
                'name': f"SGO Academy: {desc_usuario or nombre_modelo}",
                'modelo_tecnico': nombre_modelo,
                'descripcion_usuario': desc_usuario,
                'tipo': tipo,
                'descripcion_desarrollador': desc_dev,
                'tiene_menu': tiene_menu,
                'ruta': ruta_texto,
                'detalles': detalles,
                'contenido': contenido,
                'estado': 'publicado',
            })

            self.write({'ultima_fila_procesada': fila_actual})

        except ValidationError:
            raise
        except Exception as e:
            _logger.exception("Error generando publicación SGO Academy")
            raise ValidationError(f"Error con OpenAI: {str(e)}")


class SgoGuiaPaso(models.Model):
    _name = 'sgo.guia.paso'
    _description = 'Pasos del Mini-Tutorial'
    _order = 'secuencia asc'

    post_id = fields.Many2one('sgo.destacados.post', string="Post Principal", ondelete='cascade')
    secuencia = fields.Integer(string="Paso #", default=1)
    nombre_paso = fields.Char(string="Título del paso", required=True)
    instruccion = fields.Html(string="Acción a realizar")
    imagen = fields.Binary(string="Captura de Pantalla")