# -*- coding: utf-8 -*-
from odoo import models, fields, api, _
from odoo.exceptions import UserError
import logging
import json
import threading

_logger = logging.getLogger(__name__)
_FOLDER_MIME = "application/vnd.google-apps.folder"


class ArchivoCentral(models.Model):
    _name = "archivo.central"
    _description = "Archivo Central Digital"
    _rec_name = "name"

    name = fields.Char(string="Nombre del Archivo", required=True)
    root_google_id = fields.Char(string="ID Carpeta Raíz en Google Drive",
                                 help="Pega aquí el ID de la carpeta raíz en Google Drive")

    workspace_tree_ids = fields.One2many(
        "archivo.central.workspace", "browser_id", string="Carpetas")
    workspace_search = fields.Char(string="Buscar carpeta")
    workspace_view_ids = fields.Many2many(
        "archivo.central.workspace",
        "ac_ws_view_rel", "browser_id", "workspace_id",
        string="Carpetas (filtradas)",
        compute="_compute_workspace_view", store=False)

    selected_workspace_id = fields.Many2one(
        "archivo.central.workspace", string="Carpeta seleccionada")
    breadcrumb_json = fields.Text(compute="_compute_breadcrumb_json")
    search_query = fields.Char(string="Buscar documentos")

    entry_ids = fields.One2many("archivo.central.entry", "browser_id", string="Contenido")
    folder_view_ids = fields.Many2many(
        "archivo.central.entry",
        "ac_folder_view_rel", "browser_id", "entry_id",
        string="Subcarpetas", compute="_compute_entry_view", store=False)
    file_view_ids = fields.Many2many(
        "archivo.central.entry",
        "ac_file_view_rel", "browser_id", "entry_id",
        string="Documentos", compute="_compute_entry_view", store=False)
    sync_state = fields.Selection(
        [('sincronizando', 'Sincronizando'), ('completado', 'Completado')],
        string="Estado Sync Drive",
        default='completado',
    )
    upload_progress = fields.Char(string="Progreso de subida", default="")

    # ── Google Drive ──────────────────────────────────────────────────────────

    def _get_drive(self):
        from .drive import GoogleDriveAuth
        return GoogleDriveAuth.get_instance().get_drive()

    def _list_all(self, q):
        drive = self._get_drive()
        params = {
            "q": q,
            "supportsAllDrives": True,
            "includeItemsFromAllDrives": True,
            "corpora": "allDrives",
        }
        out = []
        file_list = drive.ListFile(params)
        while True:
            out.extend(file_list.GetList())
            token = getattr(file_list, "metadata", {}).get("nextPageToken")
            if not token:
                break
            params["pageToken"] = token
            file_list = drive.ListFile(params)
        return out

    def _list_children(self, folder_id):
        return self._list_all(f"'{folder_id}' in parents and trashed=false")

    def _list_subfolders_recursive(self, parent_id):
        q = f"'{parent_id}' in parents and mimeType='{_FOLDER_MIME}' and trashed=false"
        children = self._list_all(q)
        out = []
        for c in children:
            node = {
                "id": c.get("id"),
                "name": c.get("title") or c.get("name") or "Sin nombre",
                "parent_id": parent_id,
            }
            if node["id"]:
                out.append(node)
                out.extend(self._list_subfolders_recursive(node["id"]))
        return out

    # ── Sync ──────────────────────────────────────────────────────────────────

    def _cleanup_duplicate_workspaces(self):
        """Elimina registros duplicados (mismo browser_id + folder_google_id), conservando el más antiguo."""
        self.ensure_one()
        WS = self.env["archivo.central.workspace"].sudo()
        all_ws = WS.search([
            ("browser_id", "=", self.id),
            ("folder_google_id", "!=", False),
        ], order="id asc")
        seen = {}
        dupes = []
        for ws in all_ws:
            key = ws.folder_google_id
            if key in seen:
                dupes.append(ws.id)
            else:
                seen[key] = ws.id
        if dupes:
            _logger.info("Eliminando %d workspace(s) duplicados en archivo.central %s", len(dupes), self.id)
            # Si el selected apunta a un duplicado, resetearlo antes de borrar
            if self.selected_workspace_id.id in dupes:
                self.sudo().write({"selected_workspace_id": False})
            WS.browse(dupes).with_context(from_cleanup=True).unlink()

    def _sync_workspace(self):
        self.ensure_one()
        if not self.root_google_id:
            return
        root_id = self.root_google_id
        WS = self.env["archivo.central.workspace"].sudo()
        self._cleanup_duplicate_workspaces()

        root_ws = WS.search([
            ("browser_id", "=", self.id),
            ("folder_google_id", "=", root_id),
        ], limit=1)
        if not root_ws:
            try:
                drive = self._get_drive()
                gf = drive.CreateFile({"id": root_id})
                gf.FetchMetadata(fields="title")
                root_name = gf.get("title") or self.name
            except Exception:
                root_name = self.name
            root_ws = WS.create({
                "browser_id": self.id,
                "name": root_name,
                "folder_google_id": root_id,
                "parent_folder_google_id": False,
                "parent_id": False,
            })

        all_nodes = self._list_subfolders_recursive(root_id)
        drive_gids = {node["id"] for node in all_nodes} | {root_id}

        existing = {
            ws.folder_google_id: ws
            for ws in WS.search([("browser_id", "=", self.id)])
            if ws.folder_google_id
        }

        # Eliminar registros que ya no existen en Drive (huérfanos)
        for gid, ws_rec in list(existing.items()):
            if gid not in drive_gids:
                _logger.info("Eliminando workspace huérfano '%s' (%s)", ws_rec.name, gid)
                # Resetear selected_workspace_id si apunta a este registro
                if self.selected_workspace_id.id == ws_rec.id:
                    self.sudo().write({"selected_workspace_id": False})
                # Eliminar entradas kanban asociadas
                self.env["archivo.central.entry"].sudo().search([
                    ("browser_id", "=", self.id),
                    ("google_id", "=", gid),
                ]).unlink()
                ws_rec.sudo().unlink()
                del existing[gid]

        for node in all_nodes:
            rec = existing.get(node["id"])
            vals = {
                "browser_id": self.id,
                "name": node["name"],
                "folder_google_id": node["id"],
                "parent_folder_google_id": node["parent_id"],
            }
            if rec:
                diff = {k: v for k, v in vals.items() if getattr(rec, k) != v}
                if diff:
                    rec.write(diff)
            else:
                existing[node["id"]] = WS.create(vals)

        full_by_gid = dict(existing)
        full_by_gid[root_id] = root_ws
        for rec in existing.values():
            if rec.id == root_ws.id:
                continue
            parent_gid = rec.parent_folder_google_id
            parent_rec = full_by_gid.get(parent_gid) if parent_gid else root_ws
            new_parent = parent_rec.id if parent_rec else root_ws.id
            if (rec.parent_id.id if rec.parent_id else False) != new_parent:
                rec.write({"parent_id": new_parent})

        if not self.selected_workspace_id:
            self.write({"selected_workspace_id": root_ws.id})

    def _refresh_entries(self):
        self.ensure_one()
        ws = self.selected_workspace_id
        if not ws or not ws.folder_google_id:
            return
        Entry = self.env["archivo.central.entry"].sudo()
        Entry.search([
            ("browser_id", "=", self.id),
            ("parent_workspace_id", "=", ws.id),
        ]).unlink()

        for c in self._list_children(ws.folder_google_id):
            gid = c.get("id")
            title = c.get("title") or c.get("name") or "Sin nombre"
            mime = c.get("mimeType") or ""
            is_folder = mime == _FOLDER_MIME
            alt_link = (
                c.get("alternateLink")
                or (f"https://drive.google.com/file/d/{gid}/view" if gid and not is_folder else "")
            )
            linked_ws = False
            if is_folder and gid:
                lws = self.env["archivo.central.workspace"].search([
                    ("browser_id", "=", self.id),
                    ("folder_google_id", "=", gid),
                ], limit=1)
                linked_ws = lws.id if lws else False
            Entry.create({
                "browser_id": self.id,
                "parent_workspace_id": ws.id,
                "name": title,
                "google_id": gid,
                "entry_type": "folder" if is_folder else "file",
                "mime_type": mime,
                "edit_url": alt_link,
                "target_workspace_id": linked_ws,
            })

    # ── Actions ───────────────────────────────────────────────────────────────

    _ROOT_FOLDER_NAME = "ARCHIVO CENTRAL DIGITAL FINANCIERA"

    @api.model
    def _find_drive_folder_by_name(self, folder_name):
        """Search all drives for a folder with the given name; return its ID or None."""
        try:
            drive = self._get_drive()
            safe = folder_name.replace("'", "\\'")
            results = drive.ListFile({
                "q": f"mimeType='{_FOLDER_MIME}' and title='{safe}' and trashed=false",
                "supportsAllDrives": True,
                "includeItemsFromAllDrives": True,
                "corpora": "allDrives",
                "maxResults": 1,
            }).GetList()
            if results:
                return results[0].get("id")
        except Exception as e:
            _logger.warning("Drive search failed: %s", e)
        return None

    @api.model
    def action_open_main(self):
        rec = self.search([('name', '=', self._ROOT_FOLDER_NAME)], limit=1)
        if not rec:
            rec = self.create({'name': self._ROOT_FOLDER_NAME})

        rec._cleanup_duplicate_workspaces()

        if not rec.root_google_id:
            # 1. Buscar en Drive
            gid = rec._find_drive_folder_by_name(self._ROOT_FOLDER_NAME)

            # 2. Si no existe, crear la carpeta en Drive
            if not gid:
                try:
                    drive = rec._get_drive()
                    folder = drive.CreateFile({
                        'title': self._ROOT_FOLDER_NAME,
                        'mimeType': _FOLDER_MIME,
                    })
                    folder.Upload()
                    gid = folder['id']
                    _logger.info("Carpeta '%s' creada en Drive con ID %s",
                                 self._ROOT_FOLDER_NAME, gid)
                except Exception as e:
                    _logger.error("No se pudo crear la carpeta en Drive: %s", e)

            if gid:
                rec.write({'root_google_id': gid})
                try:
                    rec._sync_workspace()
                    rec._refresh_entries()
                except Exception as e:
                    _logger.warning("Auto-sync failed: %s", e)

        elif rec.root_google_id and rec.selected_workspace_id and not rec.entry_ids:
            # Ya tiene ID pero no tiene entradas cargadas → sincronizar
            try:
                rec._sync_workspace()
                rec._refresh_entries()
            except Exception as e:
                _logger.warning("Auto-sync failed: %s", e)

        return {
            'type': 'ir.actions.act_window',
            'res_model': 'archivo.central',
            'view_mode': 'form',
            'res_id': rec.id,
            'target': 'current',
        }

    def action_sync(self):
        self.ensure_one()
        if not self.root_google_id:
            raise UserError(_("Define el ID de la carpeta raíz en Google Drive primero."))
        self._cleanup_duplicate_workspaces()
        self._sync_workspace()
        self._refresh_entries()
        return {"type": "ir.actions.client", "tag": "soft_reload"}

    def action_limpiar_duplicados(self):
        """Elimina workspaces duplicados (mismo folder_google_id) sin tocar Drive."""
        self.ensure_one()
        WS = self.env["archivo.central.workspace"].sudo()
        all_ws = WS.search([
            ("browser_id", "=", self.id),
            ("folder_google_id", "!=", False),
        ], order="id asc")
        seen = {}
        selected_id = self.selected_workspace_id.id
        for ws in all_ws:
            key = ws.folder_google_id
            if key in seen:
                # Si el seleccionado era el duplicado, apuntar al original
                if ws.id == selected_id:
                    self.sudo().write({"selected_workspace_id": seen[key].id})
                if ws.child_ids:
                    ws.child_ids.write({"parent_id": seen[key].id})
                ws.unlink()
            else:
                seen[key] = ws
        return {"type": "ir.actions.client", "tag": "soft_reload"}

    def action_create_folder(self):
        self.ensure_one()
        if not self.selected_workspace_id:
            raise UserError(_("Selecciona una carpeta en el árbol primero."))
        return {
            "name": _("Nueva Carpeta"),
            "type": "ir.actions.act_window",
            "res_model": "archivo.central.create.folder",
            "view_mode": "form",
            "target": "new",
            "context": {
                "default_browser_id": self.id,
                "default_workspace_id": self.selected_workspace_id.id,
            },
        }

    def action_upload_file(self):
        self.ensure_one()
        if not self.selected_workspace_id:
            raise UserError(_("Selecciona una carpeta destino en el árbol primero."))
        view_id = self.env.ref("Matriz.view_ac_upload_form").id
        return {
            "name": _("Subir Documento"),
            "type": "ir.actions.act_window",
            "res_model": "archivo.central.upload",
            "view_mode": "form",
            "view_id": view_id,
            "target": "new",
            "context": {
                "default_browser_id": self.id,
                "default_workspace_id": self.selected_workspace_id.id,
            },
        }

    def action_upload_folder(self):
        self.ensure_one()
        if not self.selected_workspace_id:
            raise UserError(_("Selecciona una carpeta destino en el árbol primero."))
        view_id = self.env.ref("Matriz.view_ac_upload_folder_form").id
        return {
            "name": _("Subir Carpeta"),
            "type": "ir.actions.act_window",
            "res_model": "archivo.central.upload",
            "view_mode": "form",
            "view_id": view_id,
            "target": "new",
            "context": {
                "default_browser_id": self.id,
                "default_workspace_id": self.selected_workspace_id.id,
            },
        }

    def action_upload_folder_structure(self, folder_data):
        """RPC from drag & drop JS: inicia la subida en segundo plano."""
        self.ensure_one()
        ws = self.env["archivo.central.workspace"].browse(folder_data.get("workspace_id"))
        if not ws.exists() or not ws.folder_google_id:
            raise UserError(_("Carpeta de destino inválida."))

        items_data = list(folder_data.get("items", []))
        browser_id = self.id
        ws_google_id = ws.folder_google_id
        db_name = self.env.cr.dbname
        uid = self.env.uid

        self.write({"sync_state": "sincronizando", "upload_progress": ""})

        def _do_upload():
            import base64 as b64mod, tempfile, os
            from odoo.modules.registry import Registry as _Reg

            try:
                with _Reg(db_name).cursor() as cr:
                    env = api.Environment(cr, uid, {})
                    browser = env["archivo.central"].browse(browser_id)
                    gdrive = browser._get_drive()
                    folder_cache = {}

                    def ensure_path(parts, parent_gid):
                        current = parent_gid
                        for i, part in enumerate(parts):
                            key = "/".join(parts[:i + 1])
                            if key not in folder_cache:
                                f = gdrive.CreateFile({
                                    "title": part,
                                    "mimeType": _FOLDER_MIME,
                                    "parents": [{"id": current}],
                                })
                                f.Upload()
                                folder_cache[key] = f["id"]
                            current = folder_cache[key]
                        return current

                    total_items = len(items_data)
                    for idx, item in enumerate(items_data, 1):
                        parts = item["path"].split("/")
                        file_name = parts[-1]
                        short_name = (file_name or "Carpeta vacía")[:45]
                        try:
                            cr.execute(
                                "UPDATE archivo_central SET upload_progress = %s WHERE id = %s",
                                (f"{idx}/{total_items}|Preparando {short_name}", browser_id)
                            )
                            cr.commit()
                        except Exception as progress_error:
                            _logger.warning("Archivo Central progress prepare error %s: %s", idx, progress_error)

                        if item.get("empty_folder"):
                            ensure_path(parts, ws_google_id)
                            try:
                                cr.execute(
                                    "UPDATE archivo_central SET upload_progress = %s WHERE id = %s",
                                    (f"{idx}/{total_items}|{short_name}", browser_id)
                                )
                                cr.commit()
                            except Exception as progress_error:
                                _logger.warning("Archivo Central progress error %s: %s", idx, progress_error)
                            continue

                        parent_gid = (
                            ensure_path(parts[:-1], ws_google_id)
                            if len(parts) > 1 else ws_google_id
                        )

                        if file_name != ".gdkeep":
                            content = b64mod.b64decode(item["data"])
                            fd, tmp = tempfile.mkstemp()
                            try:
                                with os.fdopen(fd, "wb") as fh:
                                    fh.write(content)
                                gfile = gdrive.CreateFile({
                                    "title": file_name,
                                    "parents": [{"id": parent_gid}],
                                })
                                gfile.SetContentFile(tmp)
                                gfile.Upload()
                            finally:
                                if os.path.exists(tmp):
                                    os.remove(tmp)

                        try:
                            cr.execute(
                                "UPDATE archivo_central SET upload_progress = %s WHERE id = %s",
                                (f"{idx}/{total_items}|{short_name}", browser_id)
                            )
                            cr.commit()
                        except Exception as progress_error:
                            _logger.warning("Archivo Central progress error %s: %s", idx, progress_error)

                    try:
                        browser._sync_workspace()
                        browser._refresh_entries()
                    except Exception as refresh_error:
                        _logger.warning("Archivo Central refresh post-upload falló: %s", refresh_error)

                    cr.execute(
                        "UPDATE archivo_central SET sync_state = 'completado', upload_progress = '' WHERE id = %s",
                        (browser_id,)
                    )
                    cr.commit()
            except Exception as exc:
                _logger.error("Error crítico en subida Archivo Central: %s", exc, exc_info=True)
                try:
                    with _Reg(db_name).cursor() as cr2:
                        cr2.execute(
                            "UPDATE archivo_central SET sync_state = 'completado', upload_progress = '' WHERE id = %s",
                            (browser_id,)
                        )
                        cr2.commit()
                except Exception as reset_error:
                    _logger.warning("No se pudo resetear sync_state Archivo Central: %s", reset_error)

        t = threading.Thread(target=_do_upload, daemon=True)
        self.env.cr.postcommit.add(t.start)
        return {"background": True, "total": len(items_data)}

    def action_after_upload_refresh(self):
        self.ensure_one()
        self._sync_workspace()
        self._refresh_entries()
        return True

    # ── Computed ──────────────────────────────────────────────────────────────

    @api.depends("workspace_search", "workspace_tree_ids")
    def _compute_workspace_view(self):
        for rec in self:
            if not rec.workspace_search:
                rec.workspace_view_ids = [(6, 0, rec.workspace_tree_ids.ids)]
                continue
            matched = self.env["archivo.central.workspace"].search([
                ("browser_id", "=", rec.id),
                ("name", "ilike", rec.workspace_search),
            ])
            visible_ids = set(matched.ids)
            for f in matched:
                parent = f.parent_id
                while parent:
                    visible_ids.add(parent.id)
                    parent = parent.parent_id
            rec.workspace_view_ids = [(6, 0, list(visible_ids))]

    @api.depends("selected_workspace_id", "selected_workspace_id.parent_path",
                 "selected_workspace_id.name")
    def _compute_breadcrumb_json(self):
        WS = self.env["archivo.central.workspace"]
        for rec in self:
            ws = rec.selected_workspace_id
            if not ws:
                rec.breadcrumb_json = "[]"
                continue
            ids = [int(x) for x in (ws.parent_path or "").split("/") if x.isdigit()]
            nodes = WS.browse(ids).exists()
            id_pos = {i: p for p, i in enumerate(ids)}
            nodes = nodes.sorted(key=lambda w: id_pos.get(w.id, 10**9))
            rec.breadcrumb_json = json.dumps([
                {"id": w.id, "name": w.name}
                for w in nodes
            ])

    @api.depends("selected_workspace_id", "search_query", "entry_ids")
    def _compute_entry_view(self):
        Entry = self.env["archivo.central.entry"]
        for rec in self:
            if not rec.selected_workspace_id:
                rec.folder_view_ids = [(5, 0, 0)]
                rec.file_view_ids = [(5, 0, 0)]
                continue
            domain = [
                ("browser_id", "=", rec.id),
                ("parent_workspace_id", "=", rec.selected_workspace_id.id),
            ]
            if rec.search_query:
                domain.append(("name", "ilike", rec.search_query))
            all_e = Entry.search(domain)
            rec.folder_view_ids = [(6, 0, all_e.filtered(lambda e: e.entry_type == "folder").ids)]
            rec.file_view_ids = [(6, 0, all_e.filtered(lambda e: e.entry_type == "file").ids)]


class ArchivoCentralWorkspace(models.Model):
    _name = "archivo.central.workspace"
    _description = "Archivo Central – Carpeta"
    _parent_name = "parent_id"
    _parent_store = True
    _order = "parent_path, name"

    name = fields.Char(string="Nombre", required=True)
    browser_id = fields.Many2one(
        "archivo.central", required=True, ondelete="cascade", index=True)
    folder_google_id = fields.Char(string="Google Folder ID", index=True)
    parent_folder_google_id = fields.Char(string="Google Parent Folder ID")
    parent_id = fields.Many2one(
        "archivo.central.workspace", ondelete="cascade", index=True)
    child_ids = fields.One2many(
        "archivo.central.workspace", "parent_id", string="Subcarpetas")
    parent_path = fields.Char(index=True)
    is_selected = fields.Boolean(default=False)
    depth = fields.Integer(compute="_compute_depth", store=True)
    name_indented = fields.Char(compute="_compute_name_indented", store=False)

    @api.depends("parent_path")
    def _compute_depth(self):
        for rec in self:
            rec.depth = len([p for p in (rec.parent_path or "").strip("/").split("/") if p]) - 1

    @api.depends("name", "depth")
    def _compute_name_indented(self):
        for rec in self:
            indent = "\xa0\xa0\xa0\xa0" * (rec.depth or 0)
            rec.name_indented = indent + (rec.name or "")

    def action_select_workspace(self):
        self.ensure_one()
        browser = self.browser_id
        browser.workspace_tree_ids.write({"is_selected": False})
        self.write({"is_selected": True})
        browser.write({"selected_workspace_id": self.id, "search_query": ""})
        browser._refresh_entries()
        return False

    def action_delete_workspace(self):
        self.ensure_one()
        browser = self.browser_id
        browser_id = browser.id
        gid = self.folder_google_id
        subtree = self.search([
            ("browser_id", "=", browser_id),
            ("id", "child_of", self.id),
        ])

        # Resetear selected ANTES de cualquier unlink
        if browser.selected_workspace_id.id in subtree.ids:
            browser.sudo().write({"selected_workspace_id": False})

        # Eliminar entradas kanban asociadas
        if gid:
            self.env["archivo.central.entry"].sudo().search([
                ("browser_id", "=", browser_id),
                "|",
                ("google_id", "=", gid),
                ("parent_workspace_id", "in", subtree.ids),
            ]).unlink()

        self.sudo().unlink()

        return {"type": "ir.actions.client", "tag": "soft_reload"}

    def action_upload_to_folder(self):
        self.ensure_one()
        return {
            "name": _("Subir Documento"),
            "type": "ir.actions.act_window",
            "res_model": "archivo.central.upload",
            "view_mode": "form",
            "target": "new",
            "context": {
                "default_browser_id": self.browser_id.id,
                "default_workspace_id": self.id,
            },
        }

    def action_create_subfolder(self):
        self.ensure_one()
        return {
            "name": _("Nueva Carpeta"),
            "type": "ir.actions.act_window",
            "res_model": "archivo.central.create.folder",
            "view_mode": "form",
            "target": "new",
            "context": {
                "default_browser_id": self.browser_id.id,
                "default_workspace_id": self.id,
            },
        }


class ArchivoCentralEntry(models.Model):
    _name = "archivo.central.entry"
    _description = "Archivo Central – Documento"

    browser_id = fields.Many2one(
        "archivo.central", required=True, ondelete="cascade", index=True)
    parent_workspace_id = fields.Many2one(
        "archivo.central.workspace", ondelete="cascade")
    target_workspace_id = fields.Many2one(
        "archivo.central.workspace", ondelete="set null")
    name = fields.Char(string="Nombre")
    google_id = fields.Char(string="Google File ID")
    entry_type = fields.Selection([("folder", "Carpeta"), ("file", "Archivo")], default="file")
    mime_type = fields.Char()
    edit_url = fields.Char()

    def _find_drive_record(self):
        """Busca el registro drive con este google_id (file_id o drive_file_id)."""
        if not self.google_id:
            return None
        drive_rec = self.env["drive"].search([
            "|",
            ("file_id", "=", self.google_id),
            ("drive_file_id", "=", self.google_id),
        ], limit=1)
        return drive_rec or None

    def _get_or_create_drive_record(self):
        """Busca el registro drive por google_id; si no existe lo crea. Siempre
        asegura que el registro tenga los campos de vista correctos para mostrar
        el visor iframe (NOMBRE DEL DOCUMENTO + preview)."""
        drive_rec = self._find_drive_record()
        gid = self.google_id or ""
        embed_url = f"https://drive.google.com/file/d/{gid}/preview" if gid else ""
        edit_url = self.edit_url or (f"https://drive.google.com/file/d/{gid}/view" if gid else "")
        preview_html = (
            f'<iframe src="{embed_url}" width="100%" height="600px" '
            f'allow="autoplay" style="border:none;"></iframe>'
        ) if embed_url else ""

        if drive_rec:
            # Asegura que el registro existente muestre el visor correcto:
            # invisible_todo=False → muestra "NOMBRE DEL DOCUMENTO" + preview
            # invisible_todo=True  → muestra "Documento en Drive" + Fecha/Rango (incorrecto)
            update = {}
            if drive_rec.file_data_visible:
                update["file_data_visible"] = False
            if drive_rec.invisible_todo:
                update["invisible_todo"] = False
            if drive_rec.view_mode != "preview":
                update["view_mode"] = "preview"
            if preview_html and drive_rec.preview_html != preview_html:
                update["preview_html"] = preview_html
            if update:
                drive_rec.sudo().write(update)
        elif gid:
            drive_rec = self.env["drive"].sudo().create({
                "file_name": self.name or "Sin nombre",
                "file_id": gid,
                "drive_file_id": gid,
                "file_url": edit_url,
                "preview_url": embed_url,
                "preview_html": preview_html,
                "file_data_visible": False,
                "invisible_todo": False,
                "view_mode": "preview",
            })
        return drive_rec

    def action_open_file(self):
        self.ensure_one()
        drive_rec = self._get_or_create_drive_record()
        if not drive_rec:
            raise UserError(_("No se pudo obtener el documento de Drive."))
        view_id = self.env.ref("Matriz.view_drive_form").id
        return {
            "name": self.name,
            "type": "ir.actions.act_window",
            "res_model": "drive",
            "res_id": drive_rec.id,
            "view_mode": "form",
            "view_id": view_id,
            "target": "new",
            "context": {"create": False, "edit": True},
        }

    def action_share_entry(self):
        self.ensure_one()
        drive_rec = self._get_or_create_drive_record()
        if not drive_rec:
            raise UserError(_("No se puede compartir: archivo no vinculado a Drive."))
        return drive_rec.drive_share()

    def action_download_pdf_entry(self):
        self.ensure_one()
        drive_rec = self._get_or_create_drive_record()
        if drive_rec:
            return drive_rec.download_from_drive()
        # Descarga directa usando google_id
        if not self.google_id:
            raise UserError(_("No hay ID de Google Drive para este archivo."))
        gdrive = self.browser_id._get_drive()
        gfile = gdrive.CreateFile({"id": self.google_id})
        gfile.FetchMetadata()
        mime = gfile.get("mimeType", "")
        return self.env["drive"]._drive_direct_download_action(
            self.google_id,
            self.name or gfile.get("title", "archivo"),
            fmt="pdf",
            mime_type=mime,
        )

    def action_download_word_entry(self):
        self.ensure_one()
        drive_rec = self._get_or_create_drive_record()
        if drive_rec:
            return drive_rec.download_from_drive_as_word()
        if not self.google_id:
            raise UserError(_("No hay ID de Google Drive para este archivo."))
        import base64 as b64mod, os, uuid
        from docx import Document as DocxDocument
        from PyPDF2 import PdfReader
        gdrive = self.browser_id._get_drive()
        gfile = gdrive.CreateFile({"id": self.google_id})
        gfile.FetchMetadata()
        mime = gfile.get("mimeType", "")
        word_mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        word_mime_types = {
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "application/msword",
        }
        clean = os.path.splitext(self.name or gfile.get("title", "archivo"))[0]
        final_name = f"{clean}.docx"

        if mime == "application/vnd.google-apps.document" or mime in word_mime_types:
            return self.env["drive"]._drive_direct_download_action(
                self.google_id,
                final_name,
                fmt="word",
                mime_type=mime,
            )

        tmp_orig = os.path.join("/tmp", f"{uuid.uuid4().hex}_orig")
        tmp_docx = os.path.join("/tmp", f"{uuid.uuid4().hex}.docx")
        try:
            gfile.GetContentFile(tmp_orig)
            doc = DocxDocument()
            text_extracted = False
            try:
                with open(tmp_orig, "rb") as f:
                    for page in PdfReader(f).pages:
                        for line in (page.extract_text() or "").splitlines():
                            if line.strip():
                                doc.add_paragraph(line.strip())
                                text_extracted = True
            except Exception:
                pass
            if not text_extracted:
                doc.add_paragraph("No se pudo extraer texto del archivo original.")
            doc.save(tmp_docx)
            with open(tmp_docx, "rb") as f:
                data = b64mod.b64encode(f.read()).decode()
        finally:
            for tmp in (tmp_orig, tmp_docx):
                if os.path.exists(tmp):
                    os.remove(tmp)

        att = self.env["ir.attachment"]._drive_create_temp_download({
            "name": final_name,
            "type": "binary",
            "datas": data,
            "mimetype": word_mime,
            "res_model": self._name,
            "res_id": self.id,
        })
        return {"type": "ir.actions.act_url", "url": f"/web/content/{att.id}?download=true", "target": "self"}

    def action_download_folder(self):
        self.ensure_one()
        import base64 as b64mod, io, os, uuid, zipfile
        from concurrent.futures import ThreadPoolExecutor, as_completed
        ws = self.target_workspace_id
        if not ws or not ws.folder_google_id:
            raise UserError(_("La carpeta no tiene ID de Google Drive. Sincroniza primero."))

        folder_name = ws.name or self.name or "carpeta"
        gdrive = self.browser_id._get_drive()

        def _export_info(mime):
            if "spreadsheet" in mime:
                return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", ".xlsx"
            if "document" in mime:
                return "application/vnd.openxmlformats-officedocument.wordprocessingml.document", ".docx"
            if "presentation" in mime:
                return "application/vnd.openxmlformats-officedocument.presentationml.presentation", ".pptx"
            return None, ""

        # Fase 1: recolectar metadatos del árbol completo (sin descargar)
        pending = []
        dir_entries = []  # rutas de carpetas para preservar estructura aunque estén vacías

        def _collect_meta(folder_id, zip_path):
            q = f"'{folder_id}' in parents and trashed=false"
            try:
                items = gdrive.ListFile({"q": q}).GetList()
            except Exception as e:
                _logger.warning("Error listando carpeta %s: %s", folder_id, e)
                return
            if not items:
                dir_entries.append(zip_path + "/")
                return
            for item in items:
                item_id = item.get("id")
                item_name = (item.get("title") or item.get("name") or "archivo").strip()
                item_mime = item.get("mimeType", "")
                if not item_id:
                    continue
                if item_mime == _FOLDER_MIME:
                    dir_entries.append(f"{zip_path}/{item_name}/")
                    _collect_meta(item_id, f"{zip_path}/{item_name}")
                else:
                    export_mime, ext = _export_info(item_mime)
                    safe_name = item_name if (not ext or item_name.endswith(ext)) else f"{item_name}{ext}"
                    pending.append((item_id, export_mime, f"{zip_path}/{safe_name}"))

        _collect_meta(ws.folder_google_id, folder_name)

        # Fase 2: descargas en paralelo
        def _download_one(task):
            item_id, export_mime, zip_entry = task
            tmp = os.path.join("/tmp", f"{uuid.uuid4().hex}")
            try:
                gfile = gdrive.CreateFile({"id": item_id})
                if export_mime:
                    gfile.GetContentFile(tmp, mimetype=export_mime)
                else:
                    gfile.GetContentFile(tmp)
                with open(tmp, "rb") as f:
                    return zip_entry, f.read()
            except Exception as e:
                _logger.warning("Error descargando '%s': %s", zip_entry, e)
                return None, None
            finally:
                if os.path.exists(tmp):
                    os.remove(tmp)

        downloaded = []
        with ThreadPoolExecutor(max_workers=5) as pool:
            futures = {pool.submit(_download_one, t): t for t in pending}
            for fut in as_completed(futures):
                zip_entry, data = fut.result()
                if zip_entry and data is not None:
                    downloaded.append((zip_entry, data))

        # Fase 3: construir ZIP (primero carpetas vacías, luego archivos)
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
            for dir_path in dir_entries:
                zf.writestr(dir_path, "")
            for zip_entry, data in downloaded:
                zf.writestr(zip_entry, data)

        zip_data = b64mod.b64encode(buf.getvalue()).decode()
        att = self.env["ir.attachment"]._drive_create_temp_download({
            "name": f"{folder_name}.zip",
            "type": "binary",
            "datas": zip_data,
            "mimetype": "application/zip",
            "res_model": self._name,
            "res_id": self.id,
        })
        return {"type": "ir.actions.act_url", "url": f"/web/content/{att.id}?download=true", "target": "self"}

    def action_open_folder(self):
        self.ensure_one()
        if not self.target_workspace_id:
            raise UserError(_("Carpeta no sincronizada. Usa 'Sincronizar' para actualizar."))
        browser = self.browser_id
        browser.workspace_tree_ids.write({"is_selected": False})
        self.target_workspace_id.write({"is_selected": True})
        browser.write({"selected_workspace_id": self.target_workspace_id.id})
        browser._refresh_entries()
        return {"type": "ir.actions.client", "tag": "soft_reload"}

    def action_upload_to_entry_folder(self):
        self.ensure_one()
        if not self.target_workspace_id:
            raise UserError(_("Carpeta no sincronizada. Usa 'Sincronizar' para actualizar."))
        return {
            "name": _("Subir Documento"),
            "type": "ir.actions.act_window",
            "res_model": "archivo.central.upload",
            "view_mode": "form",
            "target": "new",
            "context": {
                "default_browser_id": self.browser_id.id,
                "default_workspace_id": self.target_workspace_id.id,
            },
        }

    def action_create_subfolder_entry(self):
        self.ensure_one()
        if not self.target_workspace_id:
            raise UserError(_("Carpeta no sincronizada. Usa 'Sincronizar' para actualizar."))
        return {
            "name": _("Nueva Carpeta"),
            "type": "ir.actions.act_window",
            "res_model": "archivo.central.create.folder",
            "view_mode": "form",
            "target": "new",
            "context": {
                "default_browser_id": self.browser_id.id,
                "default_workspace_id": self.target_workspace_id.id,
            },
        }

    def action_delete_entry(self):
        self.ensure_one()
        browser = self.browser_id
        browser_id = browser.id
        gid = self.google_id

        # Mover a papelera en Drive
        if gid:
            try:
                drive = browser._get_drive()
                drive.CreateFile({"id": gid}).Trash()
            except Exception as e:
                _logger.warning("Could not trash %s: %s", gid, e)

        # Si es carpeta, eliminar también el workspace del sidebar
        if self.entry_type == "folder" and gid:
            ws = self.env["archivo.central.workspace"].sudo().search([
                ("browser_id", "=", browser_id),
                ("folder_google_id", "=", gid),
            ], limit=1)
            if ws:
                subtree = self.env["archivo.central.workspace"].sudo().search([
                    ("browser_id", "=", browser_id),
                    ("id", "child_of", ws.id),
                ])
                if browser.selected_workspace_id.id in subtree.ids:
                    browser.sudo().write({"selected_workspace_id": False})
                self.env["archivo.central.entry"].sudo().search([
                    ("browser_id", "=", browser_id),
                    ("parent_workspace_id", "in", subtree.ids),
                ]).unlink()
                ws.sudo().unlink()

        self.sudo().unlink()

        return {"type": "ir.actions.client", "tag": "soft_reload"}


# ── Wizards ───────────────────────────────────────────────────────────────────

class ArchivoCentralCreateFolder(models.TransientModel):
    _name = "archivo.central.create.folder"
    _description = "Archivo Central – Nueva Carpeta"

    browser_id = fields.Many2one("archivo.central", required=True, ondelete="cascade")
    workspace_id = fields.Many2one("archivo.central.workspace", required=True)
    workspace_name = fields.Char(related="workspace_id.name", readonly=True)
    name = fields.Char(string="Nombre de la carpeta", required=True)

    def action_create(self):
        self.ensure_one()
        if not self.workspace_id.folder_google_id:
            raise UserError(_("La carpeta seleccionada no tiene ID de Google Drive."))
        drive = self.browser_id._get_drive()
        f = drive.CreateFile({
            "title": self.name,
            "mimeType": _FOLDER_MIME,
            "parents": [{"id": self.workspace_id.folder_google_id}],
        })
        f.Upload()
        self.browser_id._sync_workspace()
        self.browser_id._refresh_entries()
        return {"type": "ir.actions.act_window_close"}


class ArchivoCentralUpload(models.TransientModel):
    _name = "archivo.central.upload"
    _description = "Archivo Central – Subir Documento"

    browser_id = fields.Many2one("archivo.central", required=True, ondelete="cascade")
    workspace_id = fields.Many2one("archivo.central.workspace", required=True)
    workspace_name = fields.Char(related="workspace_id.name", readonly=True)
    file_data = fields.Binary(string="Archivo", attachment=True)
    file_name = fields.Char()

    def action_upload(self):
        self.ensure_one()
        import base64 as b64mod, tempfile, os
        if not self.file_data:
            raise UserError(_("Selecciona un archivo primero."))
        if not self.workspace_id.folder_google_id:
            raise UserError(_("La carpeta seleccionada no tiene ID de Google Drive."))
        gdrive = self.browser_id._get_drive()
        parent_gid = self.workspace_id.folder_google_id
        file_name = self.file_name or "documento"
        fd, tmp = tempfile.mkstemp()
        try:
            with open(tmp, "wb") as fh:
                fh.write(b64mod.b64decode(self.file_data))
            os.close(fd)
            gfile = gdrive.CreateFile({"title": file_name, "parents": [{"id": parent_gid}]})
            gfile.SetContentFile(tmp)
            gfile.Upload()
        finally:
            if os.path.exists(tmp):
                os.remove(tmp)
        self.browser_id._sync_workspace()
        self.browser_id._refresh_entries()
        return {"type": "ir.actions.act_window_close"}
