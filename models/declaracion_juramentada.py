import os
import base64
import requests
from io import BytesIO
from openpyxl import load_workbook
from openpyxl.styles import Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from odoo import models, fields, api, _
from odoo.http import request
import tempfile
import uuid
from odoo.exceptions import UserError
from odoo.osv import expression
from docxtpl import DocxTemplate  # type: ignore
from datetime import datetime
from datetime import date, timedelta
import locale
from .decorador_firmas import seguimiento_firmas
from docx import Document  # type: ignore
from docx.api import Document as DocumentApi  # type: ignore
from googleapiclient.errors import HttpError  # Asegúrate de tener google-api-python-client
from .drive import GoogleDriveAuth
import logging

_logger = logging.getLogger(__name__)



class DeclaracionJuramentada(models.Model):
    _name = 'declaracion_juramentada'
    _description = 'Declaración Juramentada'
    _inherit = ['signature.mixin', 'mail.thread', 'mail.activity.mixin']
    _rec_name = 'name'

    def message_post(self, **kwargs):
        portal_login_id = self.env.context.get('portal_chatter_login_id')
        is_public_request = bool(
            request and hasattr(request, 'httprequest') and request.httprequest.path.startswith('/public/')
        )
        if not portal_login_id and is_public_request and len(self) == 1:
            if self.login_id:
                portal_login_id = self.login_id.id
            elif self.relacion and self.relacion.login_id:
                portal_login_id = self.relacion.login_id.id
        if portal_login_id and not kwargs.get('author_id'):
            login = self.env['login'].sudo().browse(portal_login_id)
            if login.exists():
                partner = login._get_portal_chatter_partner()
                if partner:
                    kwargs['author_id'] = partner.id
        return super(DeclaracionJuramentada, self).message_post(**kwargs)

    def _reassign_portal_tracking_author(self, since):
        portal_login_id = self.env.context.get('portal_chatter_login_id')
        is_public_request = bool(
            request and hasattr(request, 'httprequest') and request.httprequest.path.startswith('/public/')
        )
        if not portal_login_id and is_public_request and len(self) == 1:
            if self.login_id:
                portal_login_id = self.login_id.id
            elif self.relacion:
                portal_login = self.env['login'].sudo().search([('contratistas', '=', self.relacion.id)], limit=1)
                portal_login_id = portal_login.id if portal_login else False
        if not portal_login_id:
            return

        login = self.env['login'].sudo().browse(portal_login_id)
        if not login.exists():
            return
        partner = login._get_portal_chatter_partner()
        if not partner:
            return

        messages = self.env['mail.message'].sudo().search([
            ('model', '=', self._name),
            ('res_id', 'in', self.ids),
            ('date', '>=', since - timedelta(seconds=3)),
            ('author_id', '!=', partner.id),
        ])
        if messages:
            messages.write({'author_id': partner.id})

    _SIGNATURE_FIELDS = {
        'firma_contratista',
        'firmado_contratista',
        'log_contratista_id',
        'firma_contratista_hash',
        'firma_declarante',
        'firmado_declarante',
        'log_declarante_id',
        'firma_declarante_hash',
    }
    _GENERATED_DOCUMENT_FIELDS = {
        'declaracion_juramentada_y_soportes_doc',
        'declaracion_juramentada_y_soportes_doc_name',
    }

    relacion = fields.Many2one('cuentas_de_cobros', string='Relación')
    login_id = fields.Many2one('login', string='Login', related='relacion.login_id')
    arl = fields.Char(string='ARL')
    

    # --- ENCABEZADO Y CONSECUTIVO ---
    consecutivo = fields.Char(string='Consecutivo')
    
    # --- INFORMACIÓN GENERAL DEL CONTRATISTA ---
    name = fields.Char(string='Nombre y Apellidos', required=True)
    nombre_apellidos = fields.Char(string='Nombre y Apellidos')
    
    # Checkboxes de Tipo de Documento (Guardamos la "X")
    sel_cc_nit = fields.Char(string='Marca CC/NIT')
    sel_ce = fields.Char(string='Marca CE')
    sel_pa = fields.Char(string='Marca PA')
    
    numero_documento = fields.Char(string='Número de Documento', related="relacion.documento_de_identificacion")
    ciudad = fields.Char(string='Ciudad', default='Cali')

    # --- FECHAS ELABORACIÓN DOCUMENTO EQUIVALENTE ---
    fecha_dia = fields.Char(string='Fecha Elab. Día')
    fecha_mes = fields.Char(string='Fecha Elab. Mes')
    fecha_anio = fields.Char(string='Fecha Elab. Año')

    # --- FECHAS DE ACTIVIDAD (DESDE - HASTA) ---
    desde_dia = fields.Char(string='Desde Día')
    desde_mes = fields.Char(string='Desde Mes')
    desde_anio = fields.Char(string='Desde Año')
    
    hasta_dia = fields.Char(string='Hasta Día')
    hasta_mes = fields.Char(string='Hasta Mes')
    hasta_anio = fields.Char(string='Hasta Año')

    # --- DATOS DE CONTACTO ---

    # Datos de Contacto (Campos con fondo gris)
    direccion_residencia = fields.Char(string='Dirección')
    ciudad_residencia2 = fields.Many2one(comodel_name='municipio', string='Ciudad Residencia')
    ciudad_contacto2 = fields.Many2one(comodel_name='municipio',string='Ciudad Contacto')
    telefono_fijo = fields.Char(string='Teléfono Fijo')
    celular = fields.Char(string='Celular')
    email = fields.Char(string='Email')
    riesgo_arl = fields.Char(string='Riesgo ARL', default='1')

    # --- CLASIFICACIÓN (SERVICIO O BIEN) ---
    servicio_check = fields.Char(string='Check Servicio')
    profesional_check = fields.Char(string='Check Profesional')
    tecnico_check = fields.Char(string='Check Técnico')
    bien_check = fields.Char(string='Check Bien')

    # --- PENSIONADO Y DECLARANTE (SI/NO) ---
    # Se guardan como Char para poner la "X" en el reporte
    pensionado_si = fields.Char(string='Pensionado SI')
    pensionado_no = fields.Char(string='Pensionado NO')
    
    declarante_si = fields.Char(string='Declarante SI')
    declarante_no = fields.Char(string='Declarante NO')

    _EXCLUSIVE_MARK_FIELDS = (
        ('servicio_check', 'bien_check', 'Servicio', 'Bien'),
        ('profesional_check', 'tecnico_check', 'Profesional', 'Técnico'),
        ('pensionado_si', 'pensionado_no', 'Pensionado: Sí', 'Pensionado: No'),
        ('declarante_si', 'declarante_no', 'Declarante de renta: Sí', 'Declarante de renta: No'),
        ('salud_prepagada_si', 'salud_prepagada_no', 'Salud prepagada: Sí', 'Salud prepagada: No'),
        ('seguros_salud_si', 'seguros_salud_no', 'Seguros de salud: Sí', 'Seguros de salud: No'),
        ('intereses_vivienda_si', 'intereses_vivienda_no', 'Intereses vivienda: Sí', 'Intereses vivienda: No'),
    )
    _EXCLUSIVE_MARK_GROUPS = (
        (
            ('sel_cc_nit', 'CC/NIT'),
            ('sel_ce', 'CE'),
            ('sel_pa', 'PA'),
        ),
        (
            ('check_cop', 'COP'),
            ('check_usd', 'USD'),
            ('check_eur', 'EUR'),
        ),
    )

    @staticmethod
    def _marca_activa(value):
        if value in (False, None):
            return False
        return str(value).strip().lower() not in ('', '0', 'false', 'no', 'n')

    def _normalizar_marcas_excluyentes(self, values):
        """Conserva una sola marca por grupo sin bloquear el formulario."""
        for field_a, field_b, _label_a, _label_b in self._EXCLUSIVE_MARK_FIELDS:
            active_a = self._marca_activa(values.get(field_a))
            active_b = self._marca_activa(values.get(field_b))
            if active_a:
                values[field_b] = False
            elif active_b:
                values[field_a] = False

        for group in self._EXCLUSIVE_MARK_GROUPS:
            marked_fields = [
                field_name
                for field_name, _label in group
                if self._marca_activa(values.get(field_name))
            ]
            if marked_fields:
                selected_field = marked_fields[0]
                for field_name, _label in group:
                    if field_name != selected_field:
                        values[field_name] = False
        return values

    # --- DETALLE ACTIVIDAD Y VALORES ---
    detalle_actividad = fields.Text(string='Detalle de la actividad')
    
    
    valor_numeros = fields.Monetary(string='Valor en Números', related='relacion.valor_cuenta')
    currency_id = fields.Many2one('res.currency', string='Moneda')
    valor_letras = fields.Char(string='Valor en Letras')

    # Monedas
    check_cop = fields.Char(string='COP Check')
    check_usd = fields.Char(string='USD Check')
    check_eur = fields.Char(string='EUR Check')


    # --- DECLARACIÓN JURAMENTADA FINAL (TEXTO INFERIOR) ---
    name_declarante = fields.Char(string='Nombre Declarante Abajo')
    doc_declarante = fields.Char(string='Documento Declarante Abajo')
    
    # Relación con Municipios (Select del XML: lugar_expedicion2)
    lugar_expedicion2 = fields.Many2one('municipio', string='Lugar de Expedición (Declarante)')
    
    # --- DEDUCCIONES (SI/NO) ---
    salud_prepagada_si = fields.Char(string='Salud Prepagada SI')
    salud_prepagada_no = fields.Char(string='Salud Prepagada NO')
    
    seguros_salud_si = fields.Char(string='Seguros Salud SI')
    seguros_salud_no = fields.Char(string='Seguros Salud NO')
    
    intereses_vivienda_si = fields.Char(string='Intereses Vivienda SI')
    intereses_vivienda_no = fields.Char(string='Intereses Vivienda NO')

    otras_rentas_exentas = fields.Char(string='Otras Rentas Exentas')

    # --- CIUDAD Y FECHA DE FIRMA ---
    dia_expedicion = fields.Char(string='Día Firma')
    ciudad_expedicion2 = fields.Many2one('municipio', string='Ciudad Firma')

    # --- TABLA DE DEPENDIENTES (One2many) ---
    dependientes_ids = fields.One2many(
        'declaracion.dependiente', 
        'declaracion_id', 
        string='Dependientes'
    )

            # Lista de meses
    MES_SELECTION = [
        ('01', 'Enero'),
        ('02', 'Febrero'),
        ('03', 'Marzo'),
        ('04', 'Abril'),
        ('05', 'Mayo'),
        ('06', 'Junio'),
        ('07', 'Julio'),
        ('08', 'Agosto'),
        ('09', 'Septiembre'),
        ('10', 'Octubre'),
        ('11', 'Noviembre'),
        ('12', 'Diciembre'),
    ]


    mes_expedicion = fields.Selection(
        MES_SELECTION,
        string='Mes',
        default=lambda self: date.today().strftime('%m'),
    )

    anio = fields.Integer(
        string='Año',
        default=lambda self: date.today().year,
    )

    anio_char = fields.Char(
        string='Año',
        compute='_compute_anio_char',
    )

    dias_2 = fields.Char(string="a los")

    @api.depends('anio')
    def _compute_anio_char(self):
        for rec in self:
            rec.anio_char = str(rec.anio or '')

    # Campo combinado (solo lectura opcional)
    mes_anio = fields.Char(
        string='Mes / Año',
        compute='_compute_mes_anio',
    )

    @api.depends('mes_expedicion', 'anio')
    def _compute_mes_anio(self):
        for rec in self:
            nombre_mes = dict(self.MES_SELECTION).get(rec.mes_expedicion, '')
            rec.mes_anio = f"{nombre_mes} {rec.anio}" if rec.mes_expedicion and rec.anio else ''

    def _sincronizar_periodo_desde_cuenta(self):
        meses = dict(self.MES_SELECTION)
        for record in self:
            cuenta = record.relacion
            if not cuenta:
                continue

            vals = {}
            if cuenta.inicio_actividad:
                inicio = fields.Date.to_date(cuenta.inicio_actividad)
                vals.update({
                    'desde_dia': str(inicio.day),
                    'desde_mes': meses.get(inicio.strftime('%m'), ''),
                    'desde_anio': str(inicio.year),
                })
            if cuenta.finalizacion_actividad:
                fin = fields.Date.to_date(cuenta.finalizacion_actividad)
                vals.update({
                    'hasta_dia': str(fin.day),
                    'hasta_mes': meses.get(fin.strftime('%m'), ''),
                    'hasta_anio': str(fin.year),
                })
            if vals:
                record.with_context(skip_refirma_por_cambio=True).write(vals)

    # ==========================================
    #   MÉTODO LÓGICO DE REEMPLAZO (TEMPLATE)
    # ==========================================
    def render_informe_template(self):
        if not load_workbook:
            raise UserError("La librería openpyxl no está instalada en el servidor.")

        self._sincronizar_periodo_desde_cuenta()

        module_dir = os.path.dirname(os.path.dirname(os.path.realpath(__file__)))
        templates_dir = os.path.join(module_dir, "templates")
        solo_declaracion = bool(
            self.relacion
            and (
                self.relacion.factura_si_o_no
                or self.relacion.factura_id
                or (
                    self.relacion.numero_factura
                    and self.relacion.numero_factura.strip()
                    and self.relacion.numero_factura.strip().upper() not in ('N/A', 'NA')
                )
            )
        )
        template_filename = (
            "declaracion_juramentada_solo.xlsx"
            if solo_declaracion
            else "declaracion_juramentada_adicion.xlsx"
        )
        template_path = os.path.join(templates_dir, template_filename)
        
        if not os.path.exists(template_path):
            raise UserError(f"No se encontró la plantilla en: {template_path}")

        # Cargar libro y hoja
        workbook = load_workbook(template_path)
        sheet = workbook.active

        valor_pesos = self.valor_numeros or 0
        valor_formateado = f"$ {valor_pesos:,.0f}".replace(",", ".")

        # --- DICCIONARIO DE VARIABLES (MAPEO NORMAL) ---
        placeholders = {
            "{{consecutivo}}": str(self.consecutivo or ""),
            "{{nombre}}": str(self.name or ""),
            "{{documento}}": str(self.numero_documento or ""),
            "{{ciudad_contacto2}}": str(self.ciudad_contacto2.name if self.ciudad_contacto2 else ""),
            "{{x_cc}}": "X" if self.sel_cc_nit else "",
            "{{x_ce}}": "X" if self.sel_ce else "",
            "{{x_pa}}": "X" if self.sel_pa else "",
            "{{fecha_dia}}": str(self.dia_expedicion or ""),
            "{{fecha_mes}}": str(self.mes_expedicion or ""),
            "{{fecha_anio}}": str(self.anio_char or ""),
            "{{direccion_residencia}}": str(self.direccion_residencia or ""),
            "{{telefono_contacto}}": str(self.celular or self.telefono_fijo or ""),
            "{{email}}": str(self.email or ""),
            "{{arl}}": str(self.riesgo_arl or ""),
            "{{arl_tipo}}": str(self.arl or ""),
            "{{ciudad_residencia2}}": str(self.ciudad_residencia2.name if self.ciudad_residencia2 else ""),
            "{{dia_expedicion}}": str(self.dia_expedicion or ""),
            "{{dia_2}}": str(self.dias_2 or ""),
            "{{mes_expedicion}}": str(self.mes_expedicion),
            "{{anio_char}}": str(self.anio_char),
            "{{desde_dia}}": str(self.desde_dia or ""),
            "{{desde_mes}}": str(self.desde_mes or ""),
            "{{desde_anio}}": str(self.desde_anio or ""),
            "{{hasta_dia}}": str(self.hasta_dia or ""),
            "{{hasta_mes}}": str(self.hasta_mes or ""),
            "{{hasta_anio}}": str(self.hasta_anio or ""),
            "{{servicio_check}}": "X" if self.servicio_check else "",
            "{{check_bien}}": "X" if self.bien_check else "",
            "{{tecnico_check}}": "X" if self.tecnico_check else "",
            "{{profesional_check}}": "X" if self.profesional_check else "",
            "{{pensionado_si}}": "X" if self.pensionado_si else "",
            "{{pensionado_no}}": "X" if self.pensionado_no else "",
            "{{declarante_si}}": "X" if self.declarante_si else "",
            "{{declarante_no}}": "X" if self.declarante_no else "",
            "{{detalle_actividad}}": str(self.detalle_actividad or ""),
            "{{valor_numeros}}": valor_formateado,
            "{{valor_letras}}": str(self.valor_letras or ""),
            "{{check_cop}}": "X" if self.check_cop else "",
            "{{check_usd}}": "X" if self.check_usd else "",
            "{{check_eur}}": "X" if self.check_eur else "",
            "{{lugar_expedicion2}}": str(self.lugar_expedicion2.name if self.lugar_expedicion2 else ""),
            "{{salud_prepagada_no}}": "X" if self.salud_prepagada_no else "",
            "{{salud_prepagada_si}}": "X" if self.salud_prepagada_si else "",
            "{{seguros_salud_no}}": "X" if self.seguros_salud_no else "",
            "{{seguros_salud_si}}": "X" if self.seguros_salud_si else "",
            "{{intereses_vivienda_no}}": "X" if self.intereses_vivienda_no else "",
            "{{intereses_vivienda_si}}": "X" if self.intereses_vivienda_si else "",
            "{{otras_rentas_exentas}}": str(self.otras_rentas_exentas or ""),
            "{{ciudad_expedicion_firma}}": str(self.ciudad_expedicion2.name if self.ciudad_expedicion2 else ""),
            "{{firma_contratista_hash}}": str(self.firma_contratista_hash or ""),
            "{{firma_contratista_nombre}}": str(self.name or ""),
            "{{firma_contratista_documento}}": self._texto_documento_firma(self.numero_documento),
            "{{firma_contratista_fecha}}": self._texto_fecha_firma(self.log_contratista_id),
            '{{firma_declarante_hash}}': str(self.firma_declarante_hash or ""),
            "{{firma_declarante_nombre}}": str(self.name_declarante or self.name or ""),
            "{{firma_declarante_documento}}": self._texto_documento_firma(self.doc_declarante or self.numero_documento),
            "{{firma_declarante_fecha}}": self._texto_fecha_firma(self.log_declarante_id),
        }

        # Localizar el bloque en la plantilla seleccionada antes de reemplazar
        # sus marcadores. Las dos plantillas pueden moverlo de fila cuando se
        # agregan textos, por lo que no debe depender de una coordenada fija.
        firma_declarante_cell = None

        # Bucle de reemplazo para las variables sueltas
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value and isinstance(cell.value, str):
                    if "{{firma_declarante_nombre}}" in cell.value:
                        firma_declarante_cell = cell
                    for placeholder, value in placeholders.items():
                        if placeholder in cell.value:
                            cell.value = cell.value.replace(placeholder, str(value))

        if not firma_declarante_cell:
            raise UserError(
                "La plantilla seleccionada no contiene el bloque de firma del declarante."
            )

        self._aplicar_formato_bloque_firma(
            sheet["A38"],
            placeholders["{{firma_contratista_nombre}}"],
            placeholders["{{firma_contratista_documento}}"],
            placeholders["{{firma_contratista_fecha}}"],
            "Firma del Contratista /Proveedor",
        )
        self._aplicar_formato_bloque_firma(
            firma_declarante_cell,
            placeholders["{{firma_declarante_nombre}}"],
            placeholders["{{firma_declarante_documento}}"],
            placeholders["{{firma_declarante_fecha}}"],
            "FIRMA DE DECLARANTE:\nDOC. IDENTIDAD: {{documento}}".replace(
                "{{documento}}", placeholders["{{documento}}"]
            ),
        )

        # ==========================================
        # LLENADO DE TABLA ONE2MANY (DEPENDIENTES)
        # ==========================================
        start_row = None
        start_col = None

        # 1. Localizamos la etiqueta inicial
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value == "{{tb_parentesco}}":
                    start_row = cell.row
                    start_col = cell.column
                    break
            if start_row:
                break

        # ==========================================
        # LLENADO DE TABLA ONE2MANY (DEPENDIENTES)
        # ==========================================
        if start_row and start_col:
            current_row = start_row
            
            for dep in self.dependientes_ids:
                # 1. Parentesco (+0) - Columna A
                sheet.cell(row=current_row, column=start_col).value = dep.parentesco or ""
                
                # 2. Edad (+1) - Columna B
                sheet.cell(row=current_row, column=start_col + 1).value = dep.edad or ""
                
                # 3. No. Identificación (+2) - Columna C
                sheet.cell(row=current_row, column=start_col + 2).value = dep.identificacion or ""
                
                # 4. Nombre (+4) <--- AQUÍ ESTÁ EL TRUCO: Saltamos la columna +3 porque es parte del ID
                sheet.cell(row=current_row, column=start_col + 4).value = dep.nombre or ""
                
                # 5. Ausencia de ingreso (+7) - Esta ya te salía bien, la mantenemos
                sheet.cell(row=current_row, column=start_col + 7).value = "X" if dep.ausencia_ingreso else ""
                
                # 6. Estudios (+8) - Esta también estaba perfecta
                sheet.cell(row=current_row, column=start_col + 8).value = dep.estudios or ""
                
                # 7. Factor físico (+9) - Y esta igual
                sheet.cell(row=current_row, column=start_col + 9).value = dep.factor_fisico or ""
                
                current_row += 1 

            if not self.dependientes_ids:
                sheet.cell(row=start_row, column=start_col).value = ""

        solo_declaracion = bool(
            self.relacion
            and (
                self.relacion.factura_si_o_no
                or self.relacion.factura_id
                or (
                    self.relacion.numero_factura
                    and self.relacion.numero_factura.strip()
                    and self.relacion.numero_factura.strip().upper() not in ('N/A', 'NA')
                )
            )
        )

        # Incluye el bloque de firma y la fila de cierre inferior en ambas
        # variantes de la plantilla.
        last_row = firma_declarante_cell.row + 1

        # Mantener compacta la introducción de la declaración. La plantilla
        # combina A44:J46, por eso esas tres filas no deben conservar alturas
        # grandes o dejan un espacio visual excesivo.
        sheet.row_dimensions[42].height = 8
        sheet.row_dimensions[43].height = 18
        sheet.row_dimensions[44].height = 10
        sheet.row_dimensions[45].height = 8
        sheet.row_dimensions[46].height = 8

        # Refuerza los bordes derecho e inferior para que las líneas finales
        # también aparezcan en la vista previa y al exportar a PDF.
        thin = Side(style='thin', color='000000')
        for row_number in range(1, last_row + 1):
            cell = sheet.cell(row=row_number, column=10)
            border = cell.border
            cell.border = Border(
                left=border.left, right=thin, top=border.top,
                bottom=border.bottom, diagonal=border.diagonal,
                diagonal_direction=border.diagonal_direction,
                diagonalUp=border.diagonalUp, diagonalDown=border.diagonalDown,
                outline=border.outline, vertical=border.vertical,
                horizontal=border.horizontal,
            )
        for column_number in range(1, 11):
            cell = sheet.cell(row=last_row, column=column_number)
            border = cell.border
            cell.border = Border(
                left=border.left, right=border.right, top=border.top,
                bottom=thin, diagonal=border.diagonal,
                diagonal_direction=border.diagonal_direction,
                diagonalUp=border.diagonalUp, diagonalDown=border.diagonalDown,
                outline=border.outline, vertical=border.vertical,
                horizontal=border.horizontal,
            )

        def configurar_hoja_impresion(hoja, fila_inicial, fila_final):
            """Limita cada formato a una página legible y oculta notas internas."""
            hoja.print_area = f"A{fila_inicial}:J{fila_final}"
            for row_number in range(1, hoja.max_row + 1):
                hoja.row_dimensions[row_number].hidden = not (
                    fila_inicial <= row_number <= fila_final
                )
                for cell in hoja[row_number]:
                    # Google agrega los comentarios de Excel como páginas
                    # finales al exportar. Son ayudas internas de la plantilla,
                    # no contenido del documento entregable.
                    if cell.comment:
                        cell.comment = None
            # La plantilla conserva formato auxiliar hasta la columna Q. Aunque
            # el área de impresión termina en J, Google puede tomar esas columnas
            # al calcular el ajuste y reducir el Documento equivalente.
            for column_number in range(11, hoja.max_column + 1):
                hoja.column_dimensions[
                    get_column_letter(column_number)
                ].hidden = True
            hoja.sheet_properties.pageSetUpPr.fitToPage = True
            hoja.page_setup.orientation = hoja.ORIENTATION_PORTRAIT
            hoja.page_setup.paperSize = hoja.PAPERSIZE_A4
            hoja.page_setup.fitToWidth = 1
            hoja.page_setup.fitToHeight = 1
            hoja.page_setup.scale = None
            hoja.page_margins = PageMargins(
                left=0.25, right=0.25, top=0.30, bottom=0.30,
                header=0.10, footer=0.10,
            )
            hoja.print_options.horizontalCentered = True
            hoja.sheet_view.showGridLines = False

        if solo_declaracion:
            sheet.title = "Declaracion juramentada"
            configurar_hoja_impresion(sheet, 40, last_row)
        else:
            # El documento equivalente y la declaración compartían una hoja.
            # Al exportarla completa Google los comprimía y además incluía las
            # notas internas de las filas 70-106. Cada formato queda ahora en
            # su propia hoja/página, conservando los datos y las firmas.
            declaracion_sheet = workbook.copy_worksheet(sheet)
            sheet.title = "Documento equivalente"
            declaracion_sheet.title = "Declaracion juramentada"
            configurar_hoja_impresion(sheet, 1, 39)
            configurar_hoja_impresion(declaracion_sheet, 40, last_row)

        # El módulo puede estar instalado en una ruta de solo lectura en
        # producción. Además, un nombre fijo provoca colisiones si dos
        # trabajadores regeneran la misma declaración al mismo tiempo.
        # Guardamos el resultado en el directorio temporal del sistema con un
        # nombre único.
        with tempfile.NamedTemporaryFile(
            prefix=f"declaracion_{self.id}_",
            suffix=".xlsx",
            delete=False,
        ) as temporary_file:
            output_path = temporary_file.name
        workbook.save(output_path)
        return output_path

    def _texto_documento_firma(self, documento):
        if not documento:
            return ""
        return f"Documento de identidad: {documento}"

    def _texto_fecha_firma(self, log_firma):
        if not log_firma or not log_firma.signature_date:
            return ""
        fecha = fields.Datetime.context_timestamp(self, log_firma.signature_date)
        meses = {
            1: "enero", 2: "febrero", 3: "marzo", 4: "abril",
            5: "mayo", 6: "junio", 7: "julio", 8: "agosto",
            9: "septiembre", 10: "octubre", 11: "noviembre", 12: "diciembre",
        }
        return f"{meses.get(fecha.month, '')} {fecha.day} de {fecha.year} a las {fecha.strftime('%I:%M %p').lower()}"

    def _aplicar_formato_bloque_firma(self, cell, nombre, documento, fecha, rotulo):
        texto = "\n".join(linea for linea in [nombre, documento, fecha, rotulo] if linea)
        try:
            from openpyxl.cell.rich_text import CellRichText, TextBlock
            from openpyxl.cell.text import InlineFont
        except Exception:
            cell.value = texto
            return

        bloques = []
        if nombre:
            bloques.append(TextBlock(InlineFont(rFont="Calibri", sz=11, b=True, color="000000"), nombre))
            bloques.append("\n")
        if documento:
            bloques.append(TextBlock(InlineFont(rFont="Calibri", sz=9, b=True, color="000000"), documento))
            bloques.append("\n")
        if fecha:
            bloques.append(TextBlock(InlineFont(rFont="Calibri", sz=9, b=True, color="000000"), fecha))
            bloques.append("\n")
        if rotulo:
            bloques.append(TextBlock(InlineFont(rFont="Calibri", sz=9, b=True, color="000000"), rotulo))
        cell.value = CellRichText(bloques)

    def generar_informe_xlsx(self):
        self.ensure_one()
        _logger.info(f"--- INICIO GENERAR INFORME PARA: {self.name} ---")
        template_rendered = False

        try:
            # 1. Renderizar y crear binario
            template_rendered = self.render_informe_template()
            
            with open(template_rendered, 'rb') as f:
                file_data = base64.b64encode(f.read())

            solo_declaracion = bool(
                self.relacion
                and (
                    self.relacion.factura_si_o_no
                    or self.relacion.factura_id
                    or (
                        self.relacion.numero_factura
                        and self.relacion.numero_factura.strip()
                        and self.relacion.numero_factura.strip().upper() not in ('N/A', 'NA')
                    )
                )
            )
            prefijo_archivo = (
                "Declaracion_Juramentada"
                if solo_declaracion
                else "Documento_Equivalente_Declaracion"
            )
            filename = f"{prefijo_archivo}_{self.name.replace(' ', '_')}.xlsx"
            
            self.write({
                'declaracion_juramentada_y_soportes_doc': file_data,
                'declaracion_juramentada_y_soportes_doc_name': filename,
            })
            _logger.info("--- PASO 1: Archivo generado ---")

            # 2. Obtener la Cuenta
            cuenta = self.relacion 
            
            if cuenta:
                _logger.info(f"--- PASO 2: Cuenta ID {cuenta.id} encontrada. Buscando Acta... ---")
                
                # --- CORRECCIÓN AQUÍ ---
                # Intentamos obtener el acta directamente del campo 'contratista_id' 
                # (Según tu código 'create', ahí guardaste el ID del acta)
                acta_record = cuenta.contratista_id

                # Si el campo se llama diferente, descomenta la siguiente línea y pon el nombre correcto:
                # acta_record = cuenta.acta_entrega_id 

                # Si sigue vacía, intentamos el search como respaldo
                if not acta_record:
                    _logger.warning("--- OJO: No estaba en el campo directo. Intentando search... ---")
                    acta_record = self.env['acta_entrega_documentos'].sudo().search([
                        ('cuenta_id', '=', cuenta.id)
                    ], limit=1)

                # 3. Verificamos si apareció
                if acta_record:
                    _logger.info(f"--- PASO 3: Acta ID {acta_record.id} ENCONTRADA ---")
                    
                    # 4. Escribir archivo. En órdenes de compra el campo de
                    # documento equivalente ya contiene la orden generada por
                    # la plataforma; la declaración se guarda en su soporte
                    # independiente para no reemplazar esa orden.
                    es_orden_compra = cuenta.selecionar in ('servicios', 'bienes')
                    # Esta escritura es una regeneración interna, no una carga
                    # manual del contratista. En etapas posteriores el modelo
                    # bloquea correctamente las cargas manuales, por lo que se
                    # debe autorizar de forma explícita solo esta operación del
                    # sistema.
                    acta_para_regeneracion = acta_record.with_context(
                        allow_person_document_write=True,
                        skip_vobo_reset=True,
                        mail_notrack=True,
                        tracking_disable=True,
                    )
                    if es_orden_compra:
                        acta_para_regeneracion.write({
                            'declaracion_juramentada_y_soportes_documento': file_data,
                            'declaracion_juramentada_y_soportes_name': filename,
                        })
                    else:
                        acta_para_regeneracion.write({
                            'documento_equivalente_documento': file_data,
                            'documento_equivalente_name': filename,
                        })
                    _logger.info("--- PASO 4: Archivo escrito en Acta ---")
                    
                    # 5. Mantener la operación dentro de la transacción de la
                    # petición. Un commit aquí invalida los savepoints usados
                    # por el controlador y deja el documento sin vincular.
                    self.env.flush_all()
                    _logger.info("--- PASO 5: Cambios preparados en la transacción ---")

                    # 6. Subida al Drive
                    if es_orden_compra and acta_record.declaracion_juramentada_y_soportes_documento:
                        _logger.info("--- PASO 6: Subiendo declaración de orden a Drive... ---")
                        try:
                            upload_ok = acta_record._upload_declaracion_juramentada_y_soportes_to_drive()
                            if not upload_ok:
                                raise UserError(
                                    "No fue posible subir la declaración juramentada a Drive. "
                                    "Inténtelo nuevamente."
                                )
                            _logger.info("--- ÉXITO TOTAL: Subida Finalizada ---")
                        except Exception as e_drive:
                            if getattr(e_drive, 'pgcode', None) in ('40001', '40P01', '55P03'):
                                raise
                            _logger.exception("--- ERROR DRIVE: %s ---", e_drive)
                            if self.env.context.get('raise_document_generation_error'):
                                raise UserError(
                                    "Las fechas se actualizaron, pero no fue posible "
                                    "actualizar la declaración en Drive. Revise los "
                                    "registros del servidor e inténtelo nuevamente."
                                ) from e_drive
                    elif acta_record.documento_equivalente_documento:
                        _logger.info("--- PASO 6: Subiendo a Drive... ---")
                        try:
                            acta_record._upload_documento_equivalente_to_drive()
                            _logger.info("--- ÉXITO TOTAL: Subida Finalizada ---")
                        except Exception as e_drive:
                            if getattr(e_drive, 'pgcode', None) in ('40001', '40P01', '55P03'):
                                raise
                            _logger.exception("--- ERROR DRIVE: %s ---", e_drive)
                            if self.env.context.get('raise_document_generation_error'):
                                raise UserError(
                                    "Las fechas se actualizaron, pero no fue posible "
                                    "actualizar el documento equivalente en Drive. "
                                    "Revise los registros del servidor e inténtelo nuevamente."
                                ) from e_drive
                else:
                    _logger.error("--- ERROR: Sigue sin encontrar el Acta asociada a esta cuenta ---")
                    # Debug: Imprimir qué tiene la cuenta
                    _logger.info(f"DATOS DE LA CUENTA: contratista_id={cuenta.contratista_id}")
                    if self.env.context.get('raise_document_generation_error'):
                        raise UserError(
                            "No fue posible actualizar el documento equivalente "
                            "porque la cuenta no tiene un acta de entrega asociada."
                        )

            return True

        except Exception as e:
            if getattr(e, 'pgcode', None) in ('40001', '40P01', '55P03'):
                raise
            _logger.exception("--- ERROR GENERAL AL GENERAR DOCUMENTO EQUIVALENTE: %s ---", e)
            if self.env.context.get('raise_document_generation_error'):
                if isinstance(e, UserError):
                    raise
                raise UserError(
                    "No fue posible regenerar el documento equivalente con las "
                    "fechas corregidas. Revise los registros del servidor e "
                    "inténtelo nuevamente."
                ) from e
            return False
        finally:
            if template_rendered and os.path.exists(template_rendered):
                try:
                    os.remove(template_rendered)
                except OSError:
                    _logger.warning(
                        "No fue posible eliminar el archivo temporal de la declaración %s: %s",
                        self.id,
                        template_rendered,
                        exc_info=True,
                    )
        
    declaracion_juramentada_y_soportes_doc = fields.Binary(string='Documento equivalente')
    declaracion_juramentada_y_soportes_doc_name = fields.Char(string='Documento equivalente')

    # Campos requeridos para el Contratista

    firma_contratista = fields.Binary(string="Firma Contratista")
    firmado_contratista = fields.Boolean(string="Firmado por Contratista")
    log_contratista_id = fields.Many2one('signature.mixin', string="Log Firma Contratista")
    firma_contratista_hash = fields.Char(string="Hash Firma Contratista", readonly=True, copy=False)

    # --- CAMPOS PARA EL DECLARANTE (NUEVOS) ---
    firma_declarante = fields.Binary(string="Firma Declarante", tracking=True)
    firmado_declarante = fields.Boolean(string="Firmado por Declarante")
    log_declarante_id = fields.Many2one('signature.mixin', string="Log Firma Declarante")
    firma_declarante_hash = fields.Char(string="Hash Firma Declarante", readonly=True, copy=False)

    @api.model_create_multi
    def create(self, vals_list):
        for vals in vals_list:
            self._normalizar_marcas_excluyentes(vals)
        return super().create(vals_list)

    def write(self, vals):
        vals = dict(vals)
        self._normalizar_marcas_excluyentes(vals)
        since = fields.Datetime.now()
        result = self._write_con_seguimiento_firmas(vals)
        self._reassign_portal_tracking_author(since)
        return result

    def write(self, vals):
        vals = dict(vals)
        self._normalizar_marcas_excluyentes(vals)
        return self._write_con_seguimiento_firmas(vals)

    # Trazabilidad de firmas (ver seguimiento_firmas en decorador_firmas.py):
    # 2 slots — Contratista/Proveedor y Declarante. Llamado desde write()
    # (arriba) tras _preparar_refirma_si_cambio_contenido, que puede
    # resetear check/log/hash a False si el contenido cambió tras una firma
    # previa (ver esa función más abajo).
    @seguimiento_firmas([
        {
            'name': 'Contratista / Proveedor',
            'img': 'firma_contratista',
            'check': 'firmado_contratista',
            'log': 'log_contratista_id',
            'hash': 'firma_contratista_hash'
        },
        {
            'name': 'Declarante',
            'img': 'firma_declarante',
            'check': 'firmado_declarante',
            'log': 'log_declarante_id',
            'hash': 'firma_declarante_hash'
        }
    ])
    def _write_con_seguimiento_firmas(self, vals):
        # Usamos super() moderno para evitar el error de NameError
        return super().write(vals)

    def _preparar_refirma_si_cambio_contenido(self, vals):
        if self.env.context.get('skip_refirma_por_cambio'):
            return
        if not self._tiene_firma_activa_previa():
            return
        campos_contenido = set(vals) - self._SIGNATURE_FIELDS - self._GENERATED_DOCUMENT_FIELDS
        campos_contenido = {campo for campo in campos_contenido if campo in self._fields}
        if not campos_contenido or not self._hay_cambio_real_en_campos(vals, campos_contenido):
            return

        self._agregar_reset_firmas_a_vals(vals)

    def _tiene_firma_activa_previa(self):
        for record in self:
            if (
                record.firmado_contratista
                or record.log_contratista_id
                or record.firma_contratista_hash
                or record.firmado_declarante
                or record.log_declarante_id
                or record.firma_declarante_hash
            ):
                return True
        return False

    def _hay_cambio_real_en_campos(self, vals, campos):
        for record in self:
            for campo in campos:
                if campo == 'dependientes_ids':
                    return True
                nuevo = vals.get(campo)
                actual = record[campo]
                field = self._fields[campo]
                if field.type == 'many2one':
                    actual = actual.id or False
                elif field.type in ('one2many', 'many2many'):
                    return True
                if (nuevo or False) != (actual or False):
                    return True
        return False

    def _agregar_reset_firmas_a_vals(self, vals):
        pares = [
            ('firmado_contratista', 'log_contratista_id', 'firma_contratista_hash'),
            ('firmado_declarante', 'log_declarante_id', 'firma_declarante_hash'),
        ]
        for check, log, hash_field in pares:
            vals.setdefault(check, False)
            vals[log] = False
            vals[hash_field] = False
        self.message_post(
            body=_(
                "⚠️ Se modificó información del formulario. Las firmas anteriores quedan sin efecto y se requiere firmar nuevamente."
            ),
            subtype_xmlid="mail.mt_note",
        )

    def action_requerir_nueva_firma_por_correccion(self):
        """Las correcciones conservan las firmas electrónicas ya registradas."""
        return True


class DeclaracionDependiente(models.Model):
    _name = 'declaracion.dependiente'
    _description = 'Dependientes de la Declaración'

    declaracion_id = fields.Many2one('declaracion_juramentada', string='Declaración Padre')
    
    parentesco = fields.Char(string='Parentesco')
    edad = fields.Integer(string='Edad')
    identificacion = fields.Char(string='Identificación')
    nombre = fields.Char(string='Nombre')
    ausencia_ingreso = fields.Boolean(string='Ausencia de Ingreso') # Checkbox
    estudios = fields.Char(string='Estudios')
    factor_fisico = fields.Char(string='Factor Físico')
