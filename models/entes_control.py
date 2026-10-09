# -*- coding: utf-8 -*-
from odoo import models, fields, api
from odoo.osv import expression
from odoo.exceptions import UserError
import base64
import io
import re

try:
    from docx import Document
except ImportError:
    Document = None


class EntesControl(models.Model):
    _name = 'entes_control'
    _description = 'Entes de Control'
    _rec_name = 'nombre_proyecto'
    _order = 'vigencia desc, create_date desc'

    name = fields.Char(string='Nombre', default='Entes de Control')
    consecutivo_proyecto = fields.Char(string='Consecutivo Proyecto', index=True)
    consecutivo = fields.Char(string='Consecutivo')
    vigencia = fields.Char(string='Vigencia')
    numero_proyecto_contrato = fields.Char(string='Número de Proyecto y/o Contrato', index=True)
    nombre_proyecto = fields.Char(string='Nombre del Proyecto', index=True)
    organo_control = fields.Char(string='Órgano de Control', index=True)
    organo_grupo = fields.Selection([
        ('procuraduria', 'Procuraduría'),
        ('contraloria', 'Contraloría'),
        ('dnp', 'DNP'),
        ('fiscalia', 'Fiscalía'),
        ('otros', 'Otros'),
    ], string='Grupo órgano', compute='_compute_organo_grupo', store=True, index=True,
       help=(
           'Se calcula automáticamente a partir del campo Órgano de Control. '
           'No se debe diligenciar manualmente.'
       ))
    seccional = fields.Char(string='Seccional')
    numero_radicado = fields.Char(string='Número de radicado')
    fecha_agregado = fields.Date(
        string='Fecha agregado',
        default=fields.Date.context_today,
        readonly=True,
        copy=False,
        help='Fecha asignada automáticamente a los registros creados desde esta funcionalidad.',
    )
    etapa = fields.Char(string='Etapa')
    hechos_investigados = fields.Text(string='Hechos investigados')
    investigador_responsable = fields.Char(string='Investigador responsable')
    ultima_actuacion = fields.Char(string='Última actuación')
    link_informacion_drive = fields.Char(string='Link de acceso (Drive)')
    carpeta_direccion_juridica = fields.Char(string='Carpeta dirección jurídica')
    estado = fields.Char(string='Estado', index=True)

    fecha_actuacion = fields.Date(string='Fecha de actuación')
    fecha_hechos = fields.Date(string='Fecha de los hechos')
    delito = fields.Char(string='Delito')
    lugar_hechos = fields.Char(string='Lugar de los hechos')
    correo_contacto = fields.Char(string='Correo de contacto')

    @api.depends('organo_control')
    def _compute_organo_grupo(self):
        for rec in self:
            texto = (rec.organo_control or '').lower()
            if 'procuradur' in texto:
                rec.organo_grupo = 'procuraduria'
            elif 'contralor' in texto:
                rec.organo_grupo = 'contraloria'
            elif 'dnp' in texto:
                rec.organo_grupo = 'dnp'
            elif 'fiscal' in texto:
                rec.organo_grupo = 'fiscalia'
            else:
                rec.organo_grupo = 'otros'

    @api.model
    def _consulta_views(self):
        return [
            (self.env.ref('Matriz.view_entes_control_kanban').id, 'kanban'),
            (self.env.ref('Matriz.view_entes_control_consulta_tree').id, 'tree'),
            (self.env.ref('Matriz.view_entes_control_record_consulta_form').id, 'form'),
        ]

    @api.model
    def _action_records(self, name='Entes de Control', domain=None):
        return {
            'type': 'ir.actions.act_window',
            'name': name,
            'res_model': 'entes_control',
            'view_mode': 'kanban,tree,form',
            'views': self._consulta_views(),
            'domain': domain or [],
            'context': {'create': False, 'edit': False, 'delete': False},
        }

    @api.model
    def action_open_records(self):
        return self._action_records()

    def action_eliminar_registro(self):
        self.ensure_one()
        self.unlink()
        return {'type': 'ir.actions.client', 'tag': 'reload'}

    def action_abrir_link_informacion_drive(self):
        self.ensure_one()
        if not self.link_informacion_drive:
            raise UserError('Este registro no tiene link de información en Drive.')
        return {
            'type': 'ir.actions.act_url',
            'url': self.link_informacion_drive,
            'target': 'new',
        }


class EntesControlConsulta(models.Model):
    _name = 'entes.control.consulta'
    _description = 'Panel de consultas de Entes de Control'

    name = fields.Char(string='Nombre', default='Panel de consultas')
    busqueda_bpin = fields.Char(string='BPIN')
    busqueda_radicado = fields.Char(string='Radicado ente de control')
    busqueda_nombre = fields.Char(string='Nombre del proyecto')
    fecha_inicio = fields.Date(string='Desde')
    fecha_fin = fields.Date(string='Hasta')
    informe_adjunto = fields.Binary(string='Informe generado')
    informe_nombre = fields.Char(string='Nombre del informe')

    @api.model
    def action_open_panel(self):
        panel = self.env.ref('Matriz.entes_control_consulta_singleton', raise_if_not_found=False)
        if not panel:
            panel = self.search([], limit=1) or self.create({'name': 'Panel de consultas'})
        panel.write({
            'busqueda_bpin': False,
            'busqueda_radicado': False,
            'busqueda_nombre': False,
        })
        return {
            'type': 'ir.actions.act_window',
            'name': 'Consultas Entes de Control',
            'res_model': 'entes.control.consulta',
            'view_mode': 'form',
            'views': [(self.env.ref('Matriz.view_entes_control_consulta_form').id, 'form')],
            'view_id': self.env.ref('Matriz.view_entes_control_consulta_form').id,
            'res_id': panel.id,
            'target': 'current',
        }

    def _action_records(self, name, domain):
        return self.env['entes_control']._action_records(name=name, domain=domain)

    def _period_domain(self):
        self.ensure_one()
        domain = []
        if self.fecha_inicio:
            domain = expression.AND([domain, expression.OR([
                [('fecha_actuacion', '>=', self.fecha_inicio)],
                [('fecha_actuacion', '=', False), ('vigencia', '>=', str(self.fecha_inicio.year))],
            ])])
        if self.fecha_fin:
            domain = expression.AND([domain, expression.OR([
                [('fecha_actuacion', '<=', self.fecha_fin)],
                [('fecha_actuacion', '=', False), ('vigencia', '<=', str(self.fecha_fin.year))],
            ])])
        return domain

    def _records_for_group(self, group):
        return self.env['entes_control'].search(expression.AND([
            [('organo_grupo', '=', group)],
            self._period_domain(),
        ]))

    def action_ver_todos(self):
        return self._action_records('Todos los entes de control', self._period_domain())

    def action_buscar(self):
        self.ensure_one()
        bpin = (self.busqueda_bpin or '').strip()
        radicado = (self.busqueda_radicado or '').strip()
        nombre = (self.busqueda_nombre or '').strip()
        domain = self._period_domain()

        if bpin:
            domain = expression.AND([
                domain,
                [('consecutivo_proyecto', 'ilike', bpin)],
            ])
        if radicado:
            domain = expression.AND([
                domain,
                [('numero_radicado', 'ilike', radicado)],
            ])
        if nombre:
            domain = expression.AND([domain, [('nombre_proyecto', 'ilike', nombre)]])

        return self._action_records('Resultado de consulta', domain)

    def action_filtrar_procuraduria(self):
        return self._action_records('Procuraduría', expression.AND([[('organo_grupo', '=', 'procuraduria')], self._period_domain()]))

    def action_filtrar_contraloria(self):
        return self._action_records('Contraloría', expression.AND([[('organo_grupo', '=', 'contraloria')], self._period_domain()]))

    def action_filtrar_dnp(self):
        return self._action_records('DNP', expression.AND([[('organo_grupo', '=', 'dnp')], self._period_domain()]))

    def action_filtrar_fiscalia(self):
        return self._action_records('Fiscalía', expression.AND([[('organo_grupo', '=', 'fiscalia')], self._period_domain()]))

    def _clean(self, value, default='No registrado'):
        text = str(value or '').strip()
        text = re.sub(r'\s+', ' ', text)
        return text or default

    def _add_key_value_table(self, document, record, rows):
        table = document.add_table(rows=0, cols=2)
        table.style = 'Table Grid'
        for label, value in rows:
            row = table.add_row().cells
            row[0].text = label
            row[1].text = self._clean(value)
        document.add_paragraph('')

    def _download_docx(self, document, filename):
        stream = io.BytesIO()
        document.save(stream)
        data = base64.b64encode(stream.getvalue())
        self.write({'informe_adjunto': data, 'informe_nombre': filename})
        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/?model=entes.control.consulta&id=%s&field=informe_adjunto&filename_field=informe_nombre&download=true' % self.id,
            'target': 'self',
        }

    def _ensure_docx(self):
        if Document is None:
            raise UserError('Falta instalar python-docx en el servidor.')

    def action_generar_fiscalia_docx(self):
        self.ensure_one()
        self._ensure_docx()
        records = self._records_for_group('fiscalia')
        document = Document()
        document.add_heading('Relación de casos Fiscalía', level=1)
        document.add_paragraph('Fuente: Entes de Control - Matriz')
        if self.fecha_inicio or self.fecha_fin:
            document.add_paragraph('Periodo: %s a %s' % (self.fecha_inicio or 'inicio', self.fecha_fin or 'fecha final'))
        document.add_paragraph('Campos no disponibles en la BD actual para varios casos: delito, lugar de los hechos, fecha de los hechos y correo de contacto.')

        for idx, rec in enumerate(records, 1):
            document.add_heading('CASO No. %s' % idx, level=2)
            self._add_key_value_table(document, rec, [
                ('SPOA / Radicado', rec.numero_radicado),
                ('Despacho asignado', rec.seccional or rec.organo_control),
                ('Descripción', rec.hechos_investigados or rec.ultima_actuacion),
                ('Delito', rec.delito),
                ('Etapa noticia criminal', rec.etapa),
                ('Lugar de los hechos', rec.lugar_hechos),
                ('Fecha de los hechos', rec.fecha_hechos),
                ('Correo', rec.correo_contacto),
                ('BPIN / Contrato', rec.numero_proyecto_contrato),
                ('Proyecto', rec.nombre_proyecto),
            ])
        return self._download_docx(document, 'relacion_casos_fiscalia.docx')

    def action_generar_procuraduria_docx(self):
        self.ensure_one()
        self._ensure_docx()
        records = self._records_for_group('procuraduria')
        document = Document()
        document.add_heading('Estado procesos Procuraduría', level=1)
        document.add_paragraph('Fuente: Entes de Control - Matriz')
        if self.fecha_inicio or self.fecha_fin:
            document.add_paragraph('Periodo: %s a %s' % (self.fecha_inicio or 'inicio', self.fecha_fin or 'fecha final'))
        document.add_paragraph('El rango por fechas se cruza con Fecha de actuación cuando existe; si no está diligenciada, se usa la vigencia del registro.')

        resumen = {}
        for rec in records:
            etapa = self._clean(rec.etapa, 'Sin etapa')
            resumen[etapa] = resumen.get(etapa, 0) + 1
        document.add_heading('Resumen por etapa', level=2)
        if resumen:
            for etapa, total in sorted(resumen.items()):
                document.add_paragraph('%s: %s' % (etapa, total), style='List Bullet')
        else:
            document.add_paragraph('No hay registros para el periodo seleccionado.')

        document.add_heading('Detalle de procesos', level=2)
        for idx, rec in enumerate(records, 1):
            document.add_heading('%s. %s' % (idx, self._clean(rec.nombre_proyecto, 'Proyecto sin nombre')), level=3)
            self._add_key_value_table(document, rec, [
                ('Vigencia', rec.vigencia),
                ('BPIN / Contrato', rec.numero_proyecto_contrato),
                ('Órgano de control', rec.organo_control),
                ('Seccional', rec.seccional),
                ('Número de radicado', rec.numero_radicado),
                ('Etapa', rec.etapa),
                ('Estado', rec.estado),
                ('Investigador responsable', rec.investigador_responsable),
                ('Última actuación', rec.ultima_actuacion),
                ('Hechos investigados', rec.hechos_investigados),
                ('Link Drive', rec.link_informacion_drive),
            ])
        return self._download_docx(document, 'estado_procesos_procuraduria.docx')


class EntesControlInformeVigencia(models.Model):
    _name = 'entes.control.informe.vigencia'
    _description = 'Vigencia disponible para informes de Entes de Control'
    _order = 'name desc'

    name = fields.Char(string='Año de vigencia', required=True, index=True)
    organo_informe = fields.Selection([
        ('todos', 'Todos'),
        ('contraloria', 'Contraloría'),
        ('procuraduria', 'Procuraduría'),
        ('fiscalia', 'Fiscalía'),
        ('dnp', 'DNP'),
    ], string='Ente de control', required=True, index=True)

    _sql_constraints = [
        (
            'entes_informe_vigencia_unique',
            'unique(name, organo_informe)',
            'La vigencia ya está registrada para este ente de control.',
        ),
    ]

    @api.model
    def _effective_group(self, record):
        organo = (record.organo_control or '').lower()
        if 'procuradur' in organo:
            return 'procuraduria'
        if 'contralor' in organo:
            return 'contraloria'
        if 'fiscal' in organo:
            return 'fiscalia'
        if 'dnp' in organo:
            return 'dnp'
        return record.organo_grupo

    @api.model
    def sync_from_entes_control(self):
        """Reconstruye las opciones con las combinaciones que existen en la matriz."""
        combinations = set()
        records = self.env['entes_control'].sudo().search([
            ('vigencia', '!=', False),
        ])
        for record in records:
            vigencia = (record.vigencia or '').strip()
            if not vigencia:
                continue
            combinations.add(('todos', vigencia))
            group = self._effective_group(record)
            if group in {'contraloria', 'procuraduria', 'fiscalia', 'dnp'}:
                combinations.add((group, vigencia))

        Option = self.sudo()
        existing = Option.search([])
        existing_by_key = {
            (option.organo_informe, option.name): option
            for option in existing
        }
        for group, vigencia in combinations:
            if (group, vigencia) not in existing_by_key:
                Option.create({
                    'organo_informe': group,
                    'name': vigencia,
                })
        obsolete = existing.filtered(
            lambda option: (option.organo_informe, option.name) not in combinations
        )
        if obsolete:
            obsolete.unlink()

    @api.model
    def name_search(self, name='', args=None, operator='ilike', limit=100):
        # Garantiza que el desplegable siempre se alimente en el momento de abrirlo,
        # incluso si el panel permanecía abierto o la acción no reconstruyó opciones.
        self.sync_from_entes_control()
        return super().name_search(
            name=name,
            args=args,
            operator=operator,
            limit=limit,
        )


class EntesControlInforme(models.Model):
    _name = 'entes.control.informe'
    _description = 'Panel de informes de Entes de Control'

    name = fields.Char(string='Nombre', default='Panel de informes')
    organo_informe = fields.Selection([
        ('todos', 'Todos'),
        ('contraloria', 'Contraloría'),
        ('procuraduria', 'Procuraduría'),
        ('fiscalia', 'Fiscalía'),
        ('dnp', 'DNP'),
    ], string='Informe', default='todos')
    vigencia_id = fields.Many2one(
        'entes.control.informe.vigencia',
        string='Año de vigencia',
        ondelete='set null',
    )
    informe_adjunto = fields.Binary(string='Informe generado')
    informe_nombre = fields.Char(string='Nombre del informe')

    _ORGANO_SEARCH_TERMS = {
        'contraloria': 'contralor',
        'procuraduria': 'procuradur',
        'fiscalia': 'fiscal',
        'dnp': 'dnp',
    }

    @api.model
    def action_open_panel(self):
        self.env['entes.control.informe.vigencia'].sync_from_entes_control()
        panel = self.env.ref('Matriz.entes_control_informe_singleton', raise_if_not_found=False)
        if panel:
            panel.write({'organo_informe': 'todos', 'vigencia_id': False})
        if not panel:
            panel = self.search([], limit=1) or self.create({'name': 'Panel de informes'})
        return {
            'type': 'ir.actions.act_window',
            'name': 'Informes Entes de Control',
            'res_model': 'entes.control.informe',
            'view_mode': 'form',
            'views': [(self.env.ref('Matriz.view_entes_control_informe_form').id, 'form')],
            'view_id': self.env.ref('Matriz.view_entes_control_informe_form').id,
            'res_id': panel.id,
            'target': 'current',
        }

    @api.onchange('organo_informe')
    def _onchange_organo_informe(self):
        self.vigencia_id = False

    def _clean(self, value, default='No registrado'):
        text = str(value or '').strip()
        text = re.sub(r'\s+', ' ', text)
        return text or default

    def _period_domain(self):
        self.ensure_one()
        domain = []
        vigencia = (self.vigencia_id.name or '').strip()
        if vigencia:
            # Los datos históricos pueden venir de Excel con espacios o textos
            # adicionales. ``ilike`` evita perder registros por esas diferencias.
            domain = expression.AND([domain, [('vigencia', 'ilike', vigencia)]])
        return domain

    def _records(self, group=None):
        domain = self._period_domain()
        if group and group != 'todos':
            organo_term = self._ORGANO_SEARCH_TERMS.get(group)
            group_domain = [('organo_grupo', '=', group)]
            if organo_term:
                # Incluye registros históricos cuya clasificación almacenada no
                # se haya recalculado, pero cuyo órgano sí identifica el ente.
                group_domain = expression.OR([
                    group_domain,
                    [('organo_control', 'ilike', organo_term)],
                ])
            domain = expression.AND([group_domain, domain])
        return self.env['entes_control'].search(domain)

    def _add_kv_table(self, document, rows):
        table = document.add_table(rows=0, cols=2)
        table.style = 'Table Grid'
        for label, value in rows:
            cells = table.add_row().cells
            cells[0].text = label
            cells[1].text = self._clean(value)
        document.add_paragraph('')

    def _period_text(self):
        partes = []
        vigencia = (self.vigencia_id.name or '').strip()
        if vigencia:
            partes.append('Vigencia %s' % vigencia)
        return '; '.join(partes) if partes else 'Todas las vigencias disponibles'

    def _download_docx(self, document, filename):
        stream = io.BytesIO()
        document.save(stream)
        self.write({
            'informe_adjunto': base64.b64encode(stream.getvalue()),
            'informe_nombre': filename,
        })
        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/?model=entes.control.informe&id=%s&field=informe_adjunto&filename_field=informe_nombre&download=true' % self.id,
            'target': 'self',
        }

    def _ensure_docx(self):
        if Document is None:
            raise UserError('Falta instalar python-docx en el servidor.')

    def action_generar_informe_docx(self):
        self.ensure_one()
        self._ensure_docx()
        dispatch = {
            'contraloria': self._build_contraloria,
            'procuraduria': self._build_procuraduria,
            'fiscalia': self._build_fiscalia,
            'dnp': self._build_dnp,
            'todos': self._build_todos,
            False: self._build_todos,
        }
        document, filename = dispatch.get(self.organo_informe, self._build_todos)()
        return self._download_docx(document, filename)

    def _base_document(self, title, records):
        document = Document()
        document.add_heading(title, level=1)
        document.add_paragraph('Fuente: Entes de Control - Matriz')
        document.add_paragraph('Periodo: %s' % self._period_text())
        document.add_paragraph('Total registros encontrados: %s' % len(records))
        return document

    def _add_summary_by_stage(self, document, records):
        summary = {}
        for rec in records:
            key = self._clean(rec.etapa, 'Sin etapa')
            summary[key] = summary.get(key, 0) + 1
        document.add_heading('Resumen por etapa', level=2)
        if not summary:
            document.add_paragraph('No hay registros para los filtros seleccionados.')
            return
        for etapa, total in sorted(summary.items()):
            document.add_paragraph('%s: %s' % (etapa, total), style='List Bullet')

    def _add_detail(self, document, records, include_fiscalia=False):
        document.add_heading('Detalle de procesos', level=2)
        for idx, rec in enumerate(records, 1):
            document.add_heading('%s. %s' % (idx, self._clean(rec.nombre_proyecto, 'Proyecto sin nombre')), level=3)
            rows = [
                ('Vigencia', rec.vigencia),
                ('BPIN / Contrato', rec.numero_proyecto_contrato),
                ('Órgano de control', rec.organo_control),
                ('Seccional / Despacho', rec.seccional),
                ('Número de radicado', rec.numero_radicado),
                ('Etapa', rec.etapa),
                ('Estado', rec.estado),
                ('Investigador responsable', rec.investigador_responsable),
                ('Última actuación', rec.ultima_actuacion),
                ('Hechos investigados', rec.hechos_investigados),
                ('Link Drive', rec.link_informacion_drive),
            ]
            if include_fiscalia:
                rows.extend([
                    ('Delito', rec.delito),
                    ('Lugar de los hechos', rec.lugar_hechos),
                    ('Fecha de los hechos', rec.fecha_hechos),
                    ('Correo', rec.correo_contacto),
                ])
            self._add_kv_table(document, rows)

    def _build_todos(self):
        records = self._records()
        document = self._base_document('Informe entes de control', records)
        document.add_paragraph('Informe generado con los filtros seleccionados. Los filtros vacíos se interpretan como todos.')
        self._add_summary_by_stage(document, records)
        self._add_detail(document, records)
        return document, 'informe_entes_control.docx'

    def _build_contraloria(self):
        records = self._records('contraloria')
        document = self._base_document('Estado procesos Contraloría', records)
        document.add_paragraph('Informe generado con base en los casos registrados para Contraloría, siguiendo la estructura general del informe base suministrado.')
        self._add_summary_by_stage(document, records)
        self._add_detail(document, records)
        return document, 'estado_procesos_contraloria.docx'

    def _build_procuraduria(self):
        records = self._records('procuraduria')
        document = self._base_document('Estado procesos Procuraduría', records)
        document.add_paragraph('Informe generado con base en los casos registrados para Procuraduría, siguiendo la estructura general del informe base suministrado.')
        self._add_summary_by_stage(document, records)
        self._add_detail(document, records)
        return document, 'estado_procesos_procuraduria.docx'

    def _build_fiscalia(self):
        records = self._records('fiscalia')
        document = self._base_document('Relación de casos Fiscalía', records)
        document.add_paragraph('Informe generado con base en los casos registrados para Fiscalía, siguiendo la estructura general del informe base suministrado.')
        document.add_paragraph('Campos recomendados para alimentar mejor este informe: delito, lugar de los hechos, fecha de los hechos y correo de contacto.')
        self._add_detail(document, records, include_fiscalia=True)
        return document, 'relacion_casos_fiscalia.docx'

    def _build_dnp(self):
        records = self._records('dnp')
        document = self._base_document('Estado procesos DNP', records)
        document.add_paragraph('Informe generado con base en los casos registrados para DNP.')
        self._add_summary_by_stage(document, records)
        self._add_detail(document, records)
        return document, 'estado_procesos_dnp.docx'
