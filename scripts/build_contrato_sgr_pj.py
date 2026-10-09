"""Convert the approved highlighted PJ contract into an SGR docxtpl template."""

import argparse
from itertools import groupby
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def replace_highlight_group(paragraph, group_number, value, keep_highlight=True):
    groups = [
        list(group) for highlighted, group in groupby(
            paragraph.runs, key=lambda run: run.font.highlight_color == WD_COLOR_INDEX.YELLOW,
        ) if highlighted
    ]
    runs = groups[group_number]
    runs[0].text = value
    if not keep_highlight:
        runs[0].font.highlight_color = None
    for run in runs[1:]:
        run.text = ''


def replace_highlighted_cell(cell, value):
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.text = ''
    cell.paragraphs[0].runs[0].text = value
    for paragraph in cell.paragraphs[1:]:
        paragraph._element.getparent().remove(paragraph._element)


def build(source, target):
    document = Document(source)
    if len(document.tables) != 2 or len(document.paragraphs) < 56:
        raise ValueError('La estructura del contrato PJ cambió; revise el mapeo antes de generarlo.')
    paragraphs = list(document.paragraphs)

    resumen = document.tables[0]
    for row, variable in {
        0: 'numero_contrato',
        4: 'nombre_contratista',
        5: 'documento_corto',
        6: 'objeto_contrato',
        7: 'clasificador_unspsc',
        8: 'valor_total_texto',
        9: 'plazo_completo',
    }.items():
        replace_highlighted_cell(resumen.rows[row].cells[1], '{{ %s }}' % variable)

    body = paragraphs[1]
    for group, variable in {
        0: 'parte_contratista',
        1: 'consideraciones_pj',
        2: 'objeto_contrato',
        3: 'valor_total_texto',
        4: 'referencia_cdp_pj',
        5: 'forma_pago',
        6: 'plazo_texto',
    }.items():
        replace_highlight_group(body, group, '{{ %s }}' % variable)
    replace_highlight_group(
        body, 7, 'El contratista se obliga a ejecutar el objeto contractual, así como a:',
        keep_highlight=False,
    )
    body.runs[33].text = ''  # The example's first consideration number.
    body.runs[109].text = ''  # The example's last consideration number.
    body.runs[110].text = ''
    body.runs[192].text = body.runs[192].text.replace(' meses, término', ', término')

    # One numbered paragraph is repeated for each line of the SGR activities field.
    for index, value in (
        (2, '{%p for actividad in actividades_lista %}'),
        (3, '{{ actividad }}'),
        (4, '{%p endfor %}'),
    ):
        paragraph = paragraphs[index]
        paragraph.runs[0].text = value
    for paragraph in paragraphs[5:46]:
        paragraph._element.getparent().remove(paragraph._element)

    clauses = paragraphs[46]
    for group, variable in {
        0: 'lugar_ejecucion_pj',
        1: 'convocatoria_pj',
        2: 'contacto_contratista_pj',
    }.items():
        replace_highlight_group(clauses, group, '{{ %s }}' % variable)
    replace_highlight_group(paragraphs[47], 0, '{{ fecha_firma_texto }}')
    replace_highlight_group(paragraphs[54], 0, '{{ representante_legal }}')
    replace_highlight_group(paragraphs[54], 1, '')
    replace_highlight_group(paragraphs[55], 0, '{{ nombre_contratista }}')

    settings = document.settings.element
    if settings.find(qn('w:updateFields')) is None:
        update_fields = OxmlElement('w:updateFields')
        update_fields.set(qn('w:val'), 'true')
        settings.append(update_fields)

    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(target)


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('source', type=Path)
    parser.add_argument('target', type=Path)
    arguments = parser.parse_args()
    build(arguments.source, arguments.target)
