"""Convert the highlighted SGR prior-study Word into a reusable template."""

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from build_contrato_sgr_pj import replace_highlight_group


def replace_paragraph(paragraph, value):
    paragraph.runs[0].text = value
    paragraph.runs[0].font.highlight_color = WD_COLOR_INDEX.YELLOW
    for run in paragraph.runs[1:]:
        run.text = ''


def build(source, target):
    document = Document(source)
    if len(document.paragraphs) < 188 or document.tables:
        raise ValueError('La estructura del estudio previo cambió; revise el mapeo antes de generarlo.')
    paragraphs = list(document.paragraphs)

    replace_paragraph(paragraphs[12], '{{ designacion_sgr }}')
    replace_highlight_group(paragraphs[17], 0, '{{ necesidad }}')
    replace_highlight_group(paragraphs[21], 0, '{{ objeto_contrato }}')

    # Keep the source's numbered-list formatting while repeating each form activity.
    for index, value in (
        (49, '{%p for actividad in actividades_lista %}'),
        (50, '{{ actividad }}'),
        (51, '{%p endfor %}'),
    ):
        replace_paragraph(paragraphs[index], value)
    for paragraph in paragraphs[52:65]:
        paragraph._element.getparent().remove(paragraph._element)

    for index, variable in (
        (81, 'perfil_requerido'),
        (83, 'experiencia_requerida'),
        (105, 'lugar_ejecucion'),
        (111, 'valor_cdp_estudio'),
        (121, 'plazo_texto'),
        (186, 'proyectado_por'),
        (187, 'revisado_por'),
    ):
        replace_highlight_group(paragraphs[index], 0, '{{ %s }}' % variable)
    paragraphs[121].runs[-1].text = paragraphs[121].runs[-1].text.replace(
        ' meses a partir', ' a partir',
    )

    payment = paragraphs[113]
    payment.runs[1].text = '{{ forma_pago }}'
    for run in payment.runs[2:]:
        run.text = ''
    paragraphs[115]._element.getparent().remove(paragraphs[115]._element)

    date_paragraph = paragraphs[175]
    date_paragraph.runs[1].text = '{{ fecha_estudio_texto }}'
    for run in date_paragraph.runs[2:]:
        run.text = ''

    settings = document.settings.element
    if settings.find(qn('w:updateFields')) is None:
        update_fields = OxmlElement('w:updateFields')
        update_fields.set(qn('w:val'), 'true')
        settings.append(update_fields)

    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(target)


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('source', type=Path)
    parser.add_argument('target', type=Path)
    arguments = parser.parse_args()
    build(arguments.source, arguments.target)
