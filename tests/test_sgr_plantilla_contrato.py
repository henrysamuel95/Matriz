from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace

from docx import Document
from docxtpl import DocxTemplate
from odoo.exceptions import UserError
from odoo.modules.module import get_module_resource
from odoo.tests.common import TransactionCase

from ..models.contrataciones_sgr.plantilla_contrato_ps import (
    ContratacionesSgrPlantillaContratoPs,
)


def _contrato(tipo):
    return SimpleNamespace(
        ensure_one=lambda: None,
        _valor_en_letras=lambda valor: 'VALOR EN LETRAS',
        select_persona=tipo,
        id_del_contrato='SGR-1',
        razon_social='Empresa Jurídica SAS',
        nombre_contratista='Persona Natural',
        contratista_proveedor_relacionado=SimpleNamespace(name='Proveedor'),
        nit='900123456-7',
        documento_identidad='12345678',
        ciudad_contratista='Cali',
        objeto_contrato='Objeto de prueba',
        justificacion_solicitud='Justificación de prueba',
        valor_total_contrato=1000000,
        forma_pago='Un pago',
        duracion='1 mes',
        actividades_especificas='Actividades jurídicas',
        actividades_especificas_natu='Actividades naturales',
        sgr_cdp_numero='CDP-1',
        sgr_cdp_fecha=date(2026, 10, 6),
        correo='correo@example.com',
        direccion_contratista='Calle 1',
        telefono_contratista='3001234567',
        r_legal='Representante de prueba',
        cedula_representante_legal='98765432',
        sgr_pj_consideraciones='Consideración uno\nConsideración dos',
        sgr_pj_unspsc='801016 - Gerencia de Proyectos',
        sgr_pj_convocatoria='Convocatoria Pública No. 1 de 2026',
        sgr_pj_lugar_ejecucion='en Cali, Valle del Cauca',
    )


class TestSgrPlantillaContrato(TransactionCase):
    def test_juridica_usa_sus_campos(self):
        contrato = _contrato('persona_juridica')

        contexto = ContratacionesSgrPlantillaContratoPs._contexto_contrato_sgr(
            contrato, 'persona_juridica',
        )

        self.assertEqual(contexto['nombre_contratista'], 'Empresa Jurídica SAS')
        self.assertEqual(contexto['documento_completo'], '900123456-7')
        self.assertEqual(contexto['actividades_especificas'], 'Actividades jurídicas')
        self.assertEqual(contexto['representante_legal'], 'Representante de prueba')
        self.assertEqual(contexto['actividades_lista'], ['Actividades jurídicas'])
        self.assertIn('98765432', contexto['parte_contratista'])
        self.assertEqual(contexto['clasificador_unspsc'], '801016 - Gerencia de Proyectos')

    def test_plantilla_pj_renderiza_sin_datos_del_ejemplo(self):
        contrato = _contrato('persona_juridica')
        contrato.actividades_especificas = 'Primera actividad\nSegunda actividad'
        contexto = ContratacionesSgrPlantillaContratoPs._contexto_contrato_sgr(
            contrato, 'persona_juridica',
        )
        plantilla = get_module_resource(
            'Matriz', 'static', 'src', 'docx', 'contrato_sgr_pj.docx',
        )
        with TemporaryDirectory() as directory:
            documento = DocxTemplate(plantilla)
            documento.render(contexto, autoescape=True)
            salida = Path(directory) / 'contrato_pj.docx'
            documento.save(salida)
            word = Document(salida)
            contenido = '\n'.join(parrafo.text for parrafo in word.paragraphs)
            contenido += '\n' + '\n'.join(
                celda.text for tabla in word.tables
                for fila in tabla.rows for celda in fila.cells
            )

        self.assertIn('Empresa Jurídica SAS', contenido)
        self.assertIn('Primera actividad', contenido)
        self.assertIn('Segunda actividad', contenido)
        self.assertIn('Convocatoria Pública No. 1 de 2026', contenido)
        self.assertNotIn('ISLAS MULTILINGÜES', contenido)
        self.assertNotIn('2026002880001', contenido)

    def test_natural_conserva_sus_campos(self):
        contrato = _contrato('persona_natural')

        contexto = ContratacionesSgrPlantillaContratoPs._contexto_contrato_sgr(
            contrato, 'persona_natural',
        )

        self.assertEqual(contexto['nombre_contratista'], 'Persona Natural')
        self.assertEqual(contexto['documento_completo'], '12345678 de Cali')
        self.assertEqual(contexto['actividades_especificas'], 'Actividades naturales')
        self.assertEqual(contexto['plazo_texto'], '1 mes')
        self.assertIn('1 mes', contexto['plazo_completo'])

    def test_juridica_requiere_nit(self):
        contrato = _contrato('persona_juridica')
        contrato.nit = False

        with self.assertRaisesRegex(UserError, 'NIT'):
            ContratacionesSgrPlantillaContratoPs._contexto_contrato_sgr(
                contrato, 'persona_juridica',
            )

    def test_juridica_requiere_campos_del_word_pj(self):
        contrato = _contrato('persona_juridica')
        contrato.sgr_pj_convocatoria = False

        with self.assertRaisesRegex(UserError, 'Convocatoria PJ'):
            ContratacionesSgrPlantillaContratoPs._contexto_contrato_sgr(
                contrato, 'persona_juridica',
            )

    def test_word_requiere_duracion_calculada(self):
        contrato = _contrato('persona_natural')
        contrato.duracion = False

        with self.assertRaisesRegex(UserError, 'Duración'):
            ContratacionesSgrPlantillaContratoPs._contexto_contrato_sgr(
                contrato, 'persona_natural',
            )
