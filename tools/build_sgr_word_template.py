"""Build the SGR natural-person Word template from the supplied contract.

Usage: python tools/build_sgr_word_template.py source.docx
"""

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.run import Run


OUTPUT = (
    Path(__file__).resolve().parents[1]
    / 'static' / 'src' / 'docx' / 'plantilla_contrato_sgr_pn.docx'
)


def runs_in_order(paragraph):
    for child in paragraph._p:
        if child.tag == qn('w:r'):
            yield Run(child, paragraph)
        elif child.tag == qn('w:hyperlink'):
            for nested in child:
                if nested.tag == qn('w:r'):
                    yield Run(nested, paragraph)


def replace_highlights(paragraph, variables):
    groups = []
    current = []
    for run in runs_in_order(paragraph):
        if run.text and run.font.highlight_color == WD_COLOR_INDEX.YELLOW:
            current.append(run)
        elif current:
            groups.append(current)
            current = []
    if current:
        groups.append(current)
    if len(groups) != len(variables):
        raise ValueError(f'Expected {len(variables)} highlighted groups; found {len(groups)}')

    for group, variable in zip(groups, variables):
        placeholder = '{{ ' + variable + ' }}'
        first = group[0]
        hyperlink = first._r.getparent()
        if hyperlink.tag == qn('w:hyperlink'):
            plain_run = OxmlElement('w:r')
            plain_text = OxmlElement('w:t')
            plain_text.text = placeholder
            plain_run.append(plain_text)
            paragraph._p.insert(paragraph._p.index(hyperlink), plain_run)
            Run(plain_run, paragraph).font.highlight_color = WD_COLOR_INDEX.YELLOW
            for run in group:
                if run._r.getparent() is not hyperlink:
                    run.text = ''
                    run.font.highlight_color = None
            paragraph._p.remove(hyperlink)
        else:
            first.text = placeholder
            first.font.highlight_color = WD_COLOR_INDEX.YELLOW
            for run in group[1:]:
                run.text = ''
                run.font.highlight_color = None


def replace_paragraph(paragraph, before, variable, after=''):
    paragraph.clear()
    paragraph.add_run(before)
    paragraph.add_run('{{ ' + variable + ' }}').font.highlight_color = WD_COLOR_INDEX.YELLOW
    if after:
        paragraph.add_run(after)


def main(source):
    document = Document(source)
    header = document.sections[0].header.tables[0]
    for paragraph, variable in zip(
        header.cell(1, 0).paragraphs,
        ('nombre_contratista', 'documento_corto'),
    ):
        replace_highlights(paragraph, [variable])

    details = document.tables[0]
    for row, variable in {
        3: 'nombre_contratista',
        4: 'documento_completo',
        5: 'objeto_contrato',
        7: 'valor_total_texto',
        8: 'plazo_completo',
    }.items():
        replace_highlights(details.cell(row, 1).paragraphs[0], [variable])

    replace_paragraph(document.paragraphs[0], 'No. ', 'numero_contrato')
    replace_highlights(document.paragraphs[2], [
        'nombre_contratista',
        'documento_completo',
        'consideraciones',
        'objeto_contrato',
        'valor_cdp_texto',
        'forma_pago',
        'plazo_texto',
        'actividades_especificas',
        'contacto_contratista',
    ])
    replace_paragraph(
        document.paragraphs[5],
        'En fe de lo anterior y para constancia de las partes, se firma en Santiago de Cali por quienes intervinieron a ',
        'fecha_firma_texto',
        '.',
    )
    replace_highlights(document.paragraphs[12], ['nombre_contratista', 'documento_completo'])
    replace_highlights(document.paragraphs[57], ['nombre_contratista'])
    replace_highlights(document.paragraphs[58], ['documento_completo'])

    used_relationships = set(document.element.xpath('.//@r:id'))
    for relation_id, relation in list(document.part.rels.items()):
        if 'hyperlink' in relation.reltype and relation_id not in used_relationships:
            del document.part.rels[relation_id]

    document.save(OUTPUT)
    print(f'Created {OUTPUT}')


if __name__ == '__main__':
    main(Path(sys.argv[1]))
